"""server/autocheck_api against the repository's synthetic comparison fixture (tests/comparison/fixtures)."""
from __future__ import annotations

import base64
import io
import json
import tempfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from autocheck_api.app import create_app
from autocheck_api.engine import RequestError, materialize

ROOT = Path(__file__).resolve().parents[2]
FIXTURES = ROOT / "tests" / "comparison" / "fixtures"
TOKEN = "test-token"


@pytest.fixture(scope="module")
def client():
    app = create_app(engine_dir=ROOT / "public" / "comparison", token=TOKEN, origins="http://localhost:8080")
    with TestClient(app) as client:
        yield client


def auth() -> dict:
    return {"Authorization": f"Bearer {TOKEN}"}


def revision(label: str, folder: str):
    """File references plus multipart parts for one fixture revision, laid out like the page sends them."""
    refs, files = [], []
    for path in sorted((FIXTURES / folder).iterdir()):
        part = f"{label}-{len(files)}"
        refs.append({"path": f"{label}/{path.name}", "part": part})
        files.append((part, (f"{label}/{path.name}", path.read_bytes(), "application/octet-stream")))
    return refs, files


def post(client, kind, request, files=(), headers=None):
    return client.post(f"/api/autocheck/{kind}", data={"request": json.dumps(request)}, files=list(files),
                       headers=auth() if headers is None else headers)


def temp_dirs() -> int:
    return sum(1 for path in Path(tempfile.gettempdir()).iterdir() if path.name.startswith("autocheck-"))


@pytest.fixture(scope="module")
def run_result(client):
    a, a_files = revision("EVT", "A")
    b, b_files = revision("DVT", "B")
    before = temp_dirs()
    response = post(client, "run", {"a": a, "b": b, "aLabel": "EVT", "bLabel": "DVT", "options": {"library": "devices: []\n"}}, a_files + b_files)
    assert response.status_code == 200, response.text
    assert temp_dirs() == before, "the request's temporary folder is removed"
    return response.json()


def test_health_reports_engine_limits_and_auth(client):
    response = client.get("/api/autocheck/health")
    assert response.status_code == 200
    body = response.json()
    assert body["ok"] and body["service"] == "autocheck-api" and body["auth"] == "token"
    assert body["engine"]["revision"] == "2026-10-27" and body["engine"]["files"] == 52
    assert body["limits"] == {"count": 1000, "file": 64 * 1024 * 1024, "total": 128 * 1024 * 1024}
    assert client.get("/api/autocheck/health", headers={"Authorization": "Bearer wrong"}).status_code == 401
    assert client.get("/api/autocheck/health", headers=auth()).status_code == 200


def test_cors_allows_only_the_configured_site(client):
    allowed = client.options("/api/autocheck/run", headers={"Origin": "http://localhost:8080", "Access-Control-Request-Method": "POST",
                                                           "Access-Control-Request-Headers": "authorization"})
    assert allowed.status_code == 200 and allowed.headers["access-control-allow-origin"] == "http://localhost:8080"
    denied = client.options("/api/autocheck/run", headers={"Origin": "http://evil.example", "Access-Control-Request-Method": "POST"})
    assert denied.status_code == 400


def test_jobs_need_the_token(client):
    assert post(client, "library", {"yaml": "devices: []\n"}, headers={}).status_code == 401
    assert post(client, "library", {"yaml": "devices: []\n"}, headers={"Authorization": "Bearer nope"}).status_code == 401
    assert post(client, "library", {"yaml": "devices: []\n"}).status_code == 200


def test_bad_requests_are_rejected_before_the_engine_runs(client):
    assert post(client, "nope", {}).status_code == 404
    assert post(client, "run", None).json() == {"error": "工作內容格式無效"}
    assert client.post("/api/autocheck/run", files=[("f0", ("x", b"x"))], headers=auth()).json() == {"error": "缺少 request 欄位"}
    assert "JSON" in client.post("/api/autocheck/run", data={"request": "{"}, headers=auth()).json()["error"]
    traversal = post(client, "run", {"a": [{"path": "../../etc/passwd", "part": "f0"}], "b": []}, [("f0", ("x", b"x"))])
    assert traversal.status_code == 400 and "上層目錄" in traversal.json()["error"]
    missing = post(client, "run", {"a": [{"path": "EVT/pstxnet.dat", "part": "nope"}], "b": []})
    assert missing.status_code == 400 and "缺少檔案內容" in missing.json()["error"]
    incomplete = post(client, "run", {"a": [{"path": "EVT/pstxnet.dat", "part": "f0"}], "b": [], "aLabel": "EVT", "bLabel": "DVT"},
                      [("f0", ("EVT/pstxnet.dat", b"NET_NAME"))])
    assert incomplete.status_code == 400 and "缺少必要 netlist" in incomplete.json()["error"]


def test_materialize_applies_the_worker_limits(tmp_path):
    parts = {f"f{i}": b"x" for i in range(1001)}
    refs = [{"path": f"{i}.dat", "part": f"f{i}"} for i in range(1001)]
    with pytest.raises(RequestError, match="輸入超過限制"):
        materialize("run", {"a": refs, "b": []}, parts, tmp_path)
    with pytest.raises(RequestError, match="重複檔案路徑"):
        materialize("run", {"a": [{"path": "x.dat", "part": "f0"}, {"path": "X.DAT", "part": "f1"}], "b": []}, parts, tmp_path)
    with pytest.raises(RequestError, match="railMap 格式"):
        materialize("run", {"a": [], "b": [], "options": {"railMap": "text"}}, parts, tmp_path)
    prepared = materialize("run", {"a": [{"path": "EVT/pstxnet.dat", "part": "f0"}], "b": [], "options": {"library": "devices: []\n"}}, parts, tmp_path)
    assert Path(prepared["a"][0]).read_bytes() == b"x" and prepared["root"] == str(tmp_path)
    assert prepared["options"]["library"] == "devices: []\n" and prepared["options"]["datasheets"] == []


def test_run_returns_the_browser_contract(run_result):
    assert run_result["meta"]["a"]["label"] == "EVT" and run_result["meta"]["b"]["label"] == "DVT"
    assert any(finding["rule"] == "MLCC_DERATING" and finding["status"] == "NEW" for finding in run_result["findings"])
    assert any(finding["status"] == "CHANGE" for finding in run_result["findings"])
    rows = run_result["rows"]
    assert rows and len({row["key"] for row in rows}) == len(rows)
    assert all(isinstance(row.get("parent"), str) and isinstance(row.get("children"), list) for row in rows)
    assert isinstance(run_result["rootcause"], dict)
    keys = [artifact["key"] for artifact in run_result["artifacts"]]
    assert len(keys) == len(set(keys))
    for key in ("pdf", "xlsx", "json", "schematic", "annotations"):
        assert key in keys, keys
    assert all("path" not in artifact for artifact in run_result["artifacts"]), "temporary paths never leave the server"
    artifacts = {artifact["key"]: artifact for artifact in run_result["artifacts"]}
    assert base64.b64decode(artifacts["schematic"]["bytes"]).startswith(b"%PDF")
    workbook = load_workbook(io.BytesIO(base64.b64decode(artifacts["xlsx"]["bytes"])))
    sheet = workbook["追蹤表"]
    head = [cell.value for cell in sheet[1]]
    assert "父項" in head and "PDF" in head
    pdf_column = head.index("PDF")
    assert all(row[pdf_column].data_type == "f" for row in sheet.iter_rows(min_row=2)), "PDF link formulas survive the safety pass"


def test_tracker_fills_replies_into_the_workbook(client, run_result):
    template = next(artifact for artifact in run_result["artifacts"] if artifact["key"] == "xlsx")
    key = run_result["rows"][0]["key"]
    response = post(client, "tracker", {"template": {"part": "t"}, "rows": run_result["rows"], "aLabel": "EVT", "bLabel": "DVT",
                                        "replies": {key: {"owner": "王小明", "note": "=1+1", "disp": "已確認"}}},
                    [("t", ("template.xlsx", base64.b64decode(template["bytes"]), "application/octet-stream"))])
    assert response.status_code == 200, response.text
    artifact = response.json()
    assert artifact["key"] == "xlsx" and artifact["name"] == "AutoCheck_tracker_EVT_vs_DVT.xlsx" and "path" not in artifact
    sheet = load_workbook(io.BytesIO(base64.b64decode(artifact["bytes"])))["追蹤表"]
    head = [cell.value for cell in sheet[1]]
    owner, note = head.index("負責人"), head.index("備註")
    rows = list(sheet.iter_rows(min_row=2))
    assert any(row[owner].value == "王小明" for row in rows)
    assert any(row[note].value == "=1+1" and row[note].data_type == "s" for row in rows), "formula-looking replies stay text"


def test_library_and_datasheet_jobs(client):
    library = post(client, "library", {"yaml": "devices: []\n"})
    assert library.status_code == 200, library.text
    body = library.json()
    assert body["errors"] == [] and isinstance(body["entries"], list) and body["entries"], "built-in library entries are returned"
    invalid = post(client, "library", {"yaml": "devices: [{name: 1}]\n"})
    assert invalid.status_code in (200, 400) and (invalid.status_code == 400 or invalid.json()["errors"]), "bad YAML entries are reported"
    from reportlab.pdfgen import canvas
    buffer = io.BytesIO()
    page = canvas.Canvas(buffer)
    page.drawString(72, 720, "TMP9999 temperature sensor")
    page.drawString(72, 700, "I2C slave address: 0x4C (7-bit)")
    page.save()
    response = post(client, "datasheet", {"file": {"path": "TMP9999.pdf", "part": "d"}, "name": "TMP9999", "mpn": "TMP9999"},
                    [("d", ("TMP9999.pdf", buffer.getvalue(), "application/pdf"))])
    assert response.status_code == 200, response.text
    body = response.json()
    # The synthetic page has no address table, so the draft is empty but complete: entry, YAML, notes and page texts.
    assert body["entry"]["mpn"] == "TMP9999" and "entry_yaml" in body
    assert isinstance(body["pageTexts"], list) and isinstance(body["notes"], list) and isinstance(body["errors"], list)
