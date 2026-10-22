"""RD replies to the tracking rows, kept in one SQLite file (stdlib, no server).

One run = one compare (tag "<A>_vs_<B>", same as static/<tag>/). Each compare re-syncs the run's row
snapshot; replies are keyed by (run, tracking key) and survive re-runs. Every field change is logged in
`history` (who is not tracked yet: no login on the test machine).

Reply fields = the yellow Excel columns: 處置 / 負責人 / 預計版次 / 備註 / 填寫日期 (auto).
"""
from __future__ import annotations

import datetime as _dt
import io
import json
import sqlite3
from contextlib import contextmanager
from pathlib import Path

REPLY_FIELDS = ["處置", "負責人", "預計版次", "備註"]
DATE_FIELD = "填寫日期"
# SQL column names (ASCII) for the reply fields
_SQL = {"處置": "disp", "負責人": "owner", "預計版次": "target_rev", "備註": "note", "填寫日期": "filled"}
_RCOLS = ", ".join(_SQL[f] for f in REPLY_FIELDS)

_SCHEMA = """
CREATE TABLE IF NOT EXISTS runs (
    run TEXT PRIMARY KEY, a_label TEXT, b_label TEXT, updated TEXT,
    pdf_url TEXT, xlsx_name TEXT, xlsx BLOB
);
CREATE TABLE IF NOT EXISTS rows (
    run TEXT, key TEXT, n INTEGER, data TEXT, PRIMARY KEY (run, key)
);
CREATE TABLE IF NOT EXISTS replies (
    run TEXT, key TEXT, disp TEXT DEFAULT '', owner TEXT DEFAULT '', target_rev TEXT DEFAULT '',
    note TEXT DEFAULT '', filled TEXT DEFAULT '', updated TEXT, PRIMARY KEY (run, key)
);
CREATE TABLE IF NOT EXISTS history (
    id INTEGER PRIMARY KEY AUTOINCREMENT, run TEXT, key TEXT, field TEXT, old TEXT, new TEXT, at TEXT
);
"""


def _now() -> str:
    return _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# first import = when this AutoCheck process (Streamlit UI) started; runs stored before it are from an earlier session
PROCESS_STARTED = _now()


def stale_note(run_info: dict, started: str | None = None) -> str:
    """Warning text when the shown run was produced before this UI was started (else "")."""
    upd = run_info.get("updated") or ""
    if not upd or upd >= (started or PROCESS_STARTED):       # same fixed format: string order = time order
        return ""
    return (f"#### 這是 {upd[:16]} 的舊結果，要看最新的請先到左側「AutoCheck 比對」按「開始檢查」\n"
            "這份結果在這次開啟 AutoCheck 之前就產生了。重跑不會洗掉已填的處置與回覆；"
            "「看圖」連到的也是那次產生的框選線路圖。")


# columns added after the first release (2026-10-20, 0p9 side-by-side viewer): an existing review.db is upgraded in
# place with ALTER TABLE (SQLite has no ADD COLUMN IF NOT EXISTS), so old runs keep their replies.
_ADDED_COLS = (("runs", "pdf_url_a", "TEXT"),      # A-side annotated PDF url ("" when A has no PDF)
               ("runs", "last_key", "TEXT"))       # last tracking key viewed in 並排瀏覽 ("回上次審到")


def _migrate(con: sqlite3.Connection) -> None:
    for table, col, typ in _ADDED_COLS:
        have = {r[1] for r in con.execute(f"PRAGMA table_info({table})")}
        if col not in have:
            con.execute(f"ALTER TABLE {table} ADD COLUMN {col} {typ} DEFAULT ''")


@contextmanager
def _db(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(str(path), timeout=15)
    try:
        con.execute("PRAGMA journal_mode=WAL")
        con.executescript(_SCHEMA)
        _migrate(con)
        yield con
        con.commit()
    finally:
        con.close()


def row_dict(row, n: int, pdf_page, pdf_page_a=None, centre=None, centre_a=None) -> dict:
    """TrackRow -> the tool-side columns of the Excel 追蹤表 (same names). "PDF 頁 A" (2026-10-20) = page of the
    same key in the A-side annotated PDF, None when A has no PDF or the part is not in A (NEW parts)."""
    return {"項次": n, "追蹤 Key": row.key, "群組 ID": row.gid, "類別": row.category, "規則": row.rule,
            "嚴重度": row.severity, "狀態": row.status, "頁": row.page, "頁名": row.page_title,
            "RefDes": ", ".join(row.refdes), "Pin": ", ".join(row.pins), "Net": ", ".join(row.nets),
            "說明 (工具)": row.title, "細節 (工具)": row.detail, "關聯 ID": ", ".join(row.related),
            "PDF 頁": pdf_page, "PDF 目的地": row.dest, "PDF 頁 A": pdf_page_a,
            # box centre [x, y] in pt from the page top-left (None = no box): explicit zoom in the 並排 viewer
            "PDF 中心": centre, "PDF 中心 A": centre_a}


def sync_run(db: Path, run: str, a_label: str, b_label: str, rows: list, pdf_pages: dict, pdf_url: str,
             pdf_pages_a: dict | None = None, pdf_url_a: str = "", centres: dict | None = None,
             centres_a: dict | None = None) -> None:
    """Replace the run's row snapshot (replies and last_key untouched)."""
    pa, ce, cea = pdf_pages_a or {}, centres or {}, centres_a or {}
    with _db(db) as con:
        con.execute("DELETE FROM rows WHERE run=?", (run,))
        con.executemany("INSERT INTO rows(run, key, n, data) VALUES (?,?,?,?)",
                        [(run, r.key, n, json.dumps(row_dict(r, n, pdf_pages.get(r.key), pa.get(r.key), ce.get(r.key), cea.get(r.key)),
                                                     ensure_ascii=False))
                         for n, r in enumerate(rows, 1)])
        con.execute("INSERT INTO runs(run, a_label, b_label, updated, pdf_url, pdf_url_a) VALUES (?,?,?,?,?,?) "
                    "ON CONFLICT(run) DO UPDATE SET a_label=excluded.a_label, b_label=excluded.b_label, "
                    "updated=excluded.updated, pdf_url=excluded.pdf_url, pdf_url_a=excluded.pdf_url_a",
                    (run, a_label, b_label, _now(), pdf_url, pdf_url_a or ""))


def last_key(db: Path, run: str) -> str:
    """Tracking key the RD last viewed in 並排瀏覽 ("" when none)."""
    if not db.exists():
        return ""
    with _db(db) as con:
        r = con.execute("SELECT last_key FROM runs WHERE run=?", (run,)).fetchone()
        return (r[0] if r else "") or ""


def set_last_key(db: Path, run: str, key: str) -> None:
    with _db(db) as con:
        con.execute("UPDATE runs SET last_key=? WHERE run=?", (key or "", run))


def store_xlsx(db: Path, run: str, xlsx: Path) -> None:
    """Keep the generated tracker as the export template (survives moving / deleting the output folder)."""
    with _db(db) as con:
        con.execute("UPDATE runs SET xlsx=?, xlsx_name=? WHERE run=?", (xlsx.read_bytes(), xlsx.name, run))


def runs(db: Path) -> list[dict]:
    if not db.exists():
        return []
    with _db(db) as con:
        cur = con.execute("SELECT run, a_label, b_label, updated, pdf_url, pdf_url_a FROM runs ORDER BY updated DESC")
        return [dict(zip(["run", "a_label", "b_label", "updated", "pdf_url", "pdf_url_a"], r)) for r in cur.fetchall()]


def replies(db: Path, run: str) -> dict[str, dict]:
    if not db.exists():
        return {}
    with _db(db) as con:
        cur = con.execute(f"SELECT key, {_RCOLS}, filled FROM replies WHERE run=?", (run,))
        return {r[0]: dict(zip(REPLY_FIELDS + [DATE_FIELD], r[1:])) for r in cur.fetchall()}


def load(db: Path, run: str) -> list[dict]:
    """Rows of a run in 項次 order, tool columns + reply columns."""
    with _db(db) as con:
        cur = con.execute("SELECT key, data FROM rows WHERE run=? ORDER BY n", (run,))
        out = [json.loads(d) for _k, d in cur.fetchall()]
    rep = replies(db, run)
    for d in out:
        d.setdefault("PDF 頁 A", None)                      # rows synced before 2026-10-20
        d.setdefault("PDF 中心", None)
        d.setdefault("PDF 中心 A", None)
        r = rep.get(d["追蹤 Key"], {})
        for f in REPLY_FIELDS + [DATE_FIELD]:
            d[f] = r.get(f, "") or ""
    return out


def save(db: Path, run: str, changes: dict[str, dict[str, str]]) -> int:
    """changes: key -> {field: new value} (REPLY_FIELDS only). Writes only real changes, stamps 填寫日期,
    logs history. Returns the number of changed cells."""
    n = 0
    today = _dt.date.today().isoformat()
    with _db(db) as con:
        for key, fields in changes.items():
            cur = con.execute(f"SELECT {_RCOLS} FROM replies WHERE run=? AND key=?", (run, key))
            old = dict(zip(REPLY_FIELDS, cur.fetchone() or [""] * len(REPLY_FIELDS)))
            diff = {f: ("" if v is None else str(v)).strip() for f, v in fields.items() if f in REPLY_FIELDS}
            diff = {f: v for f, v in diff.items() if v != (old.get(f) or "")}
            if not diff:
                continue
            con.execute("INSERT OR IGNORE INTO replies(run, key) VALUES (?,?)", (run, key))
            sets = ", ".join(f"{_SQL[f]}=?" for f in diff)
            con.execute(f"UPDATE replies SET {sets}, filled=?, updated=? WHERE run=? AND key=?",
                        [*diff.values(), today, _now(), run, key])
            con.executemany("INSERT INTO history(run, key, field, old, new, at) VALUES (?,?,?,?,?,?)",
                            [(run, key, f, old.get(f) or "", v, _now()) for f, v in diff.items()])
            n += len(diff)
    return n


def history(db: Path, run: str, limit: int = 200) -> list[dict]:
    with _db(db) as con:
        cur = con.execute("SELECT at, key, field, old, new FROM history WHERE run=? ORDER BY id DESC LIMIT ?",
                          (run, limit))
        return [dict(zip(["時間", "追蹤 Key", "欄位", "原值", "新值"], r)) for r in cur.fetchall()]


def export_xlsx(db: Path, run: str) -> tuple[str, bytes] | None:
    """The run's tracker workbook with the current replies filled into the yellow columns."""
    from openpyxl import load_workbook
    with _db(db) as con:
        r = con.execute("SELECT xlsx_name, xlsx FROM runs WHERE run=?", (run,)).fetchone()
    if not r or not r[1]:
        return None
    name, blob = r
    wb = load_workbook(io.BytesIO(blob))
    fill_replies(wb["追蹤表"], replies(db, run))
    buf = io.BytesIO()
    wb.save(buf)
    return name, buf.getvalue()


def fill_replies(ws, rep: dict[str, dict]) -> int:
    """Write replies into the 追蹤表 sheet by matching the 追蹤 Key column. Returns rows filled."""
    from .report.xlsx_tracker import COL
    n = 0
    for i in range(2, ws.max_row + 1):
        r = rep.get(ws.cell(i, COL["追蹤 Key"]).value)
        if not r:
            continue
        for f in REPLY_FIELDS + [DATE_FIELD]:
            ws.cell(i, COL[f]).value = r.get(f) or None
        n += 1
    return n

def view_hint(cfg: dict) -> str:
    a = cfg.get("annotate", {}) or {}
    if str(a.get("jump", "fitr")).lower() == "fitr":
        return ("點 Excel 的 PDF 欄或這裡的「看圖」，Chrome 會依你的視窗大小自動縮放，問題點落在畫面中間附近，"
                "任何螢幕都不用調。Chrome 要先設定「在 Chrome 中開啟 PDF」(chrome://settings/content/pdfDocuments)。")
    h = (a.get("view_hint") or "").strip()
    return (f"最佳觀看設定：{h}。在這個設定下，點 Excel 的 PDF 欄或這裡的「看圖」，問題點會落在畫面正中間；"
            "其他螢幕設定也一定在畫面內，只是會偏離中央。") if h else ""
