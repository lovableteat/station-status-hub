"""B1 power tree as boxes (owner request 2026-10-12): one layout, two renderers.

layout()      power_tree() dict (B, optionally A) -> list[Box] with (level, top, height) in text-line units.
              Pure data, no Excel / SVG. Left to right: level = column, top = first text line of the box.
render_xlsx() a worksheet: one cell per text line, box = cell borders, connectors = borders in two narrow
              columns between levels (no drawing shapes: openpyxl drops shapes on read-back, issue #2075).
render_svg()  the same coordinates as SVG for the web page / a standalone .html.

Owner decisions (2026-10-12): box = rail + V + RefDes, everything else on hover (Excel comment, SVG <title>);
loads as "負載 ×N"; EVT -> DVT colours in v1; series 0R / bead / R <= 100 already merged into one rail by
RailGroups; left to right.
Graph rules: a DNP second source with the same in / out rails is a line inside the main box, not a branch;
an IC's internal LDO / bias rail hangs under that IC's box (dashed) instead of being a fake root; a rail fed
from two places is drawn once, under the first, with "也由 ... 供電"; rails with loads but no identified source
go to a red section with the built D / R parts that join them to other rails (netlist facts, no guessing).
Change status is keyed on converter RefDes (company rule: RefDes is never re-annotated between revisions).
"""
from __future__ import annotations

from dataclasses import dataclass, field
from html import escape

from ..model import natkey

STATUS_ZH = {"new": "新增", "changed": "變更", "removed": "移除"}
LEGEND = [("", "一般 (實線框 = 有上件)"), ("dnp", "虛線灰框 = DNP / IC 內部 LDO"), ("new", "綠 = 新增"),
          ("changed", "黃 = 變更 (換料 / 上件 / rail 名)"), ("removed", "紅虛線 = 移除 (A 有 B 無)"),
          ("orphan", "紅粗框 = 有負載但來源未識別")]
ANCHOR = 2            # connectors leave / enter a box at the boundary after its 2nd text line


@dataclass
class Box:
    id: int
    kind: str                       # root / conv / bias / orphan / removed / unreached
    rail: str
    volt: float | None
    level: int = 0
    top: int = 0
    lines: list[str] = field(default_factory=list)
    hover: list[str] = field(default_factory=list)
    parent: int | None = None
    conv: str = ""
    built: bool = True
    status: str = ""                # "" / new / changed / removed
    page: int | None = None         # schematic page of the converter (link target)
    section: str = "tree"           # tree / orphan / removed

    @property
    def height(self) -> int:
        return len(self.lines)

    @property
    def dashed(self) -> bool:
        return self.kind in ("bias", "removed") or not self.built


def _v(v) -> str:
    return "? V" if v is None else f"{v:g} V"


def _a_index(pt_a: dict | None) -> dict | None:
    if not pt_a or "edges" not in pt_a:
        return None
    by_conv: dict[str, list] = {}
    for e in pt_a["edges"]:
        by_conv.setdefault(e["conv"], []).append(e)
    rails = set(pt_a.get("rails", {})) | {e["rail_in"] for e in pt_a["edges"]} | {e["rail_out"] for e in pt_a["edges"]}
    return {"by_conv": by_conv, "rails": rails}


def _edge_status(e: dict, ai: dict | None) -> tuple[str, str]:
    """(status, note) of one B edge against A, keyed on converter RefDes."""
    if ai is None:
        return "", ""
    olds = ai["by_conv"].get(e["conv"])
    if not olds:
        return "new", f"{e['conv']} 新增"
    o = next((x for x in olds if x["rail_out"] == e["rail_out"]), olds[0])
    notes = []
    if o["mpn"] != e["mpn"]:
        notes.append(f"換料 {o['mpn']} -> {e['mpn']}")
    if bool(o["built"]) != bool(e["built"]):
        notes.append(f"上件 {'Y' if o['built'] else 'DNP'} -> {'Y' if e['built'] else 'DNP'}")
    if o["rail_in"] != e["rail_in"]:
        notes.append(f"上游 rail {o['rail_in']} -> {e['rail_in']}")
    if o["rail_out"] != e["rail_out"]:
        notes.append(f"下游 rail {o['rail_out']} -> {e['rail_out']}")
    return ("changed", "; ".join(notes)) if notes else ("", "")


def layout(pt: dict, pt_a: dict | None = None, a_label: str = "A", b_label: str = "B") -> list[Box]:
    """Boxes with level / top assigned. Every converter edge of pt["edges"] appears exactly once, either as a
    box, as a second-source line inside a box, or as a "也由 ... 供電" note on the box of its output rail."""
    edges = pt.get("edges") or []
    rails = pt.get("rails") or {}
    orphans = pt.get("orphans") or []
    ai = _a_index(pt_a)
    boxes: list[Box] = []

    def new_box(**kw) -> Box:
        if kw.get("parent") is not None:                      # a child lives in its parent's section
            kw["section"] = boxes[kw["parent"]].section
        b = Box(id=len(boxes), **kw)
        boxes.append(b)
        return b

    def rinfo(r):
        return rails.get(r, {"v": None, "members": [r], "loads": [], "conns": []})

    def load_line(r, skip: str = ""):
        ri = rinfo(r)
        s = f"負載 ×{len([x for x in ri['loads'] if x != skip])}"
        if ri["conns"]:
            s += "  → " + ", ".join(ri["conns"][:3]) + (" …" if len(ri["conns"]) > 3 else "")
        return s

    def rail_hover(r):
        ri = rinfo(r)
        h = []
        if len(ri["members"]) > 1:
            h.append("同組 net: " + ", ".join(ri["members"][:12]) + (" …" if len(ri["members"]) > 12 else ""))
        if ri["loads"]:
            h.append("負載 IC: " + ", ".join(ri["loads"][:20]) + (" …" if len(ri["loads"]) > 20 else ""))
        if ri["conns"]:
            h.append("接頭: " + ", ".join(ri["conns"]))
        return h

    def tag(st):
        return f" [{STATUS_ZH[st]}]" if st else ""

    # group parallel edges (same in / out rail): main = first built, others = second-source lines
    groups: dict[tuple, list] = {}
    for e in edges:
        groups.setdefault((e["rail_in"], e["rail_out"]), []).append(e)
    for k in groups:
        groups[k].sort(key=lambda e: (not e["built"], natkey(e["conv"])))
    gkey = lambda k: (not groups[k][0]["built"], natkey(groups[k][0]["conv"]), k)  # noqa: E731
    kids: dict[str, list[tuple]] = {}
    for k in sorted(groups, key=gkey):
        kids.setdefault(k[0], []).append(k)
    # a rail fed from several places is drawn under ONE feed, chosen up front: built before DNP, then RefDes
    primary: dict[str, tuple] = {}
    for k in sorted(groups, key=gkey):
        primary.setdefault(k[1], k)

    # internal LDO / bias rails hang under the converter that makes them
    conv_rail = {}                                            # converter RefDes -> its (in, out) group
    for k in sorted(groups, key=gkey):
        for e in groups[k]:
            conv_rail.setdefault(e["conv"], k)
    by_rail_orphan = {o["rail"]: o for o in orphans}
    bias_of: dict[tuple, list] = {}
    orphan_left = []
    for o in orphans:
        k = conv_rail.get(o.get("parent_ic") or "")
        if k:
            bias_of.setdefault(k, []).append(o)
        else:
            orphan_left.append(o)
    bias_rails = {o["rail"] for os_ in bias_of.values() for o in os_}

    roots = [r["tree"] for r in pt.get("rows", []) if r.get("level") == 0 and str(r.get("kind", "")).startswith("input")]
    roots = [r for r in roots if r not in bias_rails]
    placed: set[str] = set()
    fed_by: dict[str, Box] = {}
    used_edges: set[int] = set()
    deferred: list[tuple] = []                                # (group) drawn as a note once its rail has a box

    def edge_box(k, parent: Box, level: int) -> Box:
        es = groups[k]
        main = es[0]
        st, note = _edge_status(main, ai)
        b = new_box(kind="conv", rail=k[1], volt=rinfo(k[1])["v"], level=level, parent=parent.id, conv=main["conv"],
                    built=bool(main["built"]), status=st, page=main["page"])
        b.lines = [k[1], _v(b.volt), f"{main['conv']}{'' if main['built'] else ' (DNP)'}{tag(st)}", load_line(k[1])]
        b.hover = [f"{main['conv']} {main['mpn']}  {main['kind']}  上件 {'Y' if main['built'] else 'DNP'}",
                   f"上游 {k[0]} ({_v(rinfo(k[0])['v'])})"]
        if main["en"]:
            b.hover.append("EN: " + ", ".join(main["en"]))
        if main["pg"]:
            b.hover.append("PG: " + ", ".join(main["pg"]))
        if main["via"]:
            b.hover.append("經過: " + ", ".join(main["via"]))
        if main["page"] is not None:
            b.hover.append(f"線路圖第 {main['page']} 頁")
        if note:
            b.hover.append(f"{a_label} -> {b_label}: {note}")
        for alt in es[1:]:
            ast, anote = _edge_status(alt, ai)
            b.lines.append(f"alt {alt['conv']}{'' if alt['built'] else ' (DNP)'}{tag(ast)}")
            b.hover.append(f"alt {alt['conv']} {alt['mpn']} 上件 {'Y' if alt['built'] else 'DNP'}"
                           + (f"; {a_label} -> {b_label}: {anote}" if anote else ""))
            if ast and not b.status:
                b.status = "changed"
        b.hover += rail_hover(k[1])
        for e in es:
            used_edges.add(id(e))
        return b

    def grow(rail: str, box: Box, level: int, path: frozenset):
        for k in kids.get(rail, []):
            if primary.get(k[1]) != k or k[1] in path or k[1] in placed:
                deferred.append(k)                            # second feed / loop: note, no redraw
                for e in groups[k]:
                    used_edges.add(id(e))
                continue
            placed.add(k[1])
            b = edge_box(k, box, level)
            fed_by[k[1]] = b
            for o in bias_of.get(k, []):
                bias_box(o, b, level + 1, path | {k[1]})
            grow(k[1], b, level + 1, path | {k[1]})

    def bias_box(o: dict, parent: Box | None, level: int, path: frozenset):
        r, ic = o["rail"], o.get("parent_ic", "")
        if r in placed:                                       # also the output of a (DNP) LDO drawn elsewhere
            first = fed_by.get(r)
            if first is not None:
                first.lines.append(f"也由 {ic} 內部 LDO 供電")
                first.hover.append(f"也由 {ic} 內部 LDO 供電 ({o.get('source', '')})")
            return
        placed.add(r)
        b = new_box(kind="bias", rail=r, volt=o.get("voltage"), level=level,
                    parent=parent.id if parent else None, conv=ic, built=True,
                    section="tree" if parent else "orphan")
        b.lines = [r, _v(b.volt), f"{ic} 內部 LDO" + ("" if parent else " (IC 框未畫)"), load_line(r, skip=ic)]
        b.hover = [f"工具判斷: {o.get('source', '')}"] + rail_hover(r)
        fed_by[r] = b
        grow(r, b, level + 1, path | {r})

    def orphan_lines(o: dict) -> tuple[list[str], list[str]]:
        adj = o.get("adjacent", [])
        return ([f"{a['refdes']} ↔ {a['rail']}" for a in adj],
                [f"工具判斷: {o.get('source', '')}"] + [f"{a['refdes']} {a['value']} 接到 {a['rail']}" for a in adj])

    def root_box(r: str, section: str = "tree"):
        placed.add(r)
        ri = rinfo(r)
        st = "new" if ai is not None and r not in ai["rails"] else ""
        b = new_box(kind="root", rail=r, volt=ri["v"], level=0, status=st, section=section)
        b.lines = [r, _v(b.volt), "輸入" + (" " + ", ".join(ri["conns"][:3]) if ri["conns"] else "") + tag(st),
                   load_line(r)]
        b.hover = ["電源輸入 (沒有上游轉換元件)"]
        o = by_rail_orphan.get(r)
        if o and not str(o.get("source", "")).startswith(("connector", "root input")):
            extra, hv = orphan_lines(o)                       # a root the tool could not explain: keep the facts
            b.lines += extra
            b.hover += hv
        b.hover += rail_hover(r)
        fed_by[r] = b
        grow(r, b, 1, frozenset({r}))

    for r in roots:
        if r not in placed:
            root_box(r)
    # orphans fed from a connector are inputs, not unknowns
    for o in list(orphan_left):
        if str(o.get("source", "")).startswith("connector") and o["rail"] not in placed:
            root_box(o["rail"])

    def unreached(skip_bias: bool):
        for k in sorted(groups, key=gkey):
            if id(groups[k][0]) in used_edges or k[0] in placed or (skip_bias and k[0] in bias_rails):
                continue
            placed.add(k[0])
            ri = rinfo(k[0])
            b = new_box(kind="unreached", rail=k[0], volt=ri["v"], level=0, built=False)
            b.lines = [k[0], _v(b.volt), "未接到任何輸入", load_line(k[0])]
            b.hover = ["這個 rail 只餵給轉換元件，工具串不到它的來源 (迴圈或只有 DNP 路徑)"] + rail_hover(k[0])
            fed_by[k[0]] = b
            grow(k[0], b, 1, frozenset({k[0]}))

    unreached(skip_bias=True)
    # internal-LDO rails whose IC box was never drawn (the IC only shows as a note / alt line)
    for os_ in bias_of.values():
        for o in os_:
            if o["rail"] not in placed:
                bias_box(o, None, 0, frozenset())
    unreached(skip_bias=False)

    # orphans: rails with loads, no identified source
    for o in orphan_left:
        r = o["rail"]
        if r in placed:
            continue
        placed.add(r)
        b = new_box(kind="orphan", rail=r, volt=o.get("voltage"), level=0, section="orphan")
        extra, hv = orphan_lines(o)
        b.lines = [r, _v(b.volt), "來源未識別", load_line(r)] + extra
        b.hover = hv + rail_hover(r)
        fed_by[r] = b
        grow(r, b, 1, frozenset({r}))

    # second feeds / loop edges: a note (with its own change status) on the box of the rail they feed
    for k in deferred:
        first = fed_by.get(k[1])
        if first is None:
            continue
        for e in groups[k]:
            st, note = _edge_status(e, ai)
            first.lines.append(f"也由 {k[0]} ({e['conv']}{'' if e['built'] else ' DNP'}){tag(st)} 供電")
            first.hover.append(f"也由 {k[0]} 經 {e['conv']} {e['mpn']} 上件 {'Y' if e['built'] else 'DNP'} 供電"
                               + (f"; {a_label} -> {b_label}: {note}" if note else ""))
            if st and not first.status:
                first.status = "changed"

    # A edges whose converter is gone in B
    if ai is not None:
        b_outs: dict[str, set] = {}
        for e in edges:
            b_outs.setdefault(e["conv"], set()).add(e["rail_out"])
        for c in sorted(ai["by_conv"], key=natkey):
            olds = ai["by_conv"][c]
            if c in b_outs:
                # same RefDes, fewer outputs (dual-output IC lost a rail); an equal count is a rename -> "changed"
                if len(olds) <= len(b_outs[c]):
                    continue
                olds = [e for e in olds if e["rail_out"] not in b_outs[c]]
            for e in olds:
                parent = fed_by.get(e["rail_in"])
                b = new_box(kind="removed", rail=e["rail_out"], volt=None, level=(parent.level + 1) if parent else 0,
                            parent=parent.id if parent else None, conv=c, built=bool(e["built"]), status="removed",
                            section="removed")
                b.lines = [e["rail_out"], f"{a_label} 有", f"{c} [移除]", f"上游 {e['rail_in']}"]
                b.hover = [f"{a_label} 有、{b_label} 沒有: {c} {e['mpn']} {e['kind']}",
                           f"{e['rail_in']} -> {e['rail_out']}"]
                if e.get("page") is not None:
                    b.hover.append(f"{a_label} 線路圖第 {e['page']} 頁")
    _assign_rows(boxes)
    return boxes


def _assign_rows(boxes: list[Box], gap: int = 1, section_gap: int = 3) -> None:
    """Tidy left-to-right tree: a parent shares the top line of its first child; leaves stack downwards."""
    kids: dict[int, list[Box]] = {}
    for b in boxes:
        if b.parent is not None:
            kids.setdefault(b.parent, []).append(b)

    def place(b: Box, row: int) -> int:
        b.top = row
        nxt = row
        ch = kids.get(b.id, [])
        for c in ch:
            nxt = place(c, nxt)
        return max(nxt, row + b.height + gap)

    row = 0
    for sec in ("tree", "orphan", "removed"):
        tops = [b for b in boxes if b.parent is None and b.section == sec]
        if not tops:
            continue
        if row:
            row += section_gap
        for b in tops:
            row = place(b, row)


def section_starts(boxes: list[Box]) -> dict[str, int]:
    out = {}
    for b in boxes:
        if b.parent is None:
            out[b.section] = min(out.get(b.section, b.top), b.top)
    return out


SECTION_TITLE = {"orphan": "有負載但來源未識別的 rail (紅框)：列出把它接到其他 rail 的已上件 D / R，請對線路圖確認來源",
                 "removed": "{a} 有、{b} 沒有的轉換元件 (找不到對應的上游 rail)"}


# ------------------------------------------------------------------------------------------------------- Excel
def render_xlsx(ws, boxes: list[Box], title: str, pdf_link=None, a_label: str = "A", b_label: str = "B",
                diff: bool = True) -> None:
    """pdf_link(page, text) -> Excel formula (with its leading '=') or None. Box column of level L = 2 + 3L,
    connectors in the 2 columns after it. Text line i of a box sits in row top0 + top + i."""
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    FONT = "Arial"
    fills = {"new": "E2EFDA", "changed": "FFF2CC", "removed": "F8CBAD", "orphan": "FDE9E7", "dnp": "F2F2F2"}
    line = Side(style="thin", color="404040")
    top0 = 6
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, title).font = Font(name=FONT, bold=True, size=12)
    ws.cell(2, 1, "框內：rail 名 / 電壓 / 轉換元件 RefDes (點 RefDes 跳線路圖該頁) / 負載 IC 數。滑鼠停在 rail 名上看料號、"
                  "EN / PG、同組 net、負載清單。").font = Font(name=FONT, size=9)
    lg = ws.cell(3, 1, "圖例：" + "　".join(t for _, t in LEGEND) + ("" if diff else f"　({a_label} 的 power tree 沒有算出來，"
                                                                                      "本圖沒有變更顏色)"))
    lg.font = Font(name=FONT, size=9, color="595959")
    maxlvl = max((b.level for b in boxes), default=0)
    ws.column_dimensions["A"].width = 2
    for lv in range(maxlvl + 1):
        c = 2 + 3 * lv
        ws.column_dimensions[get_column_letter(c)].width = 27
        ws.column_dimensions[get_column_letter(c + 1)].width = 2.5
        ws.column_dimensions[get_column_letter(c + 2)].width = 2.5

    starts = section_starts(boxes)
    for sec, r0 in starts.items():
        if sec in SECTION_TITLE:
            ws.cell(top0 + r0 - 2, 2, SECTION_TITLE[sec].format(a=a_label, b=b_label)).font = Font(
                name=FONT, bold=True, size=10, color="C00000")

    def add_border(cell, **kw):
        b = cell.border
        cell.border = Border(left=kw.get("left", b.left), right=kw.get("right", b.right),
                             top=kw.get("top", b.top), bottom=kw.get("bottom", b.bottom))

    by_id = {b.id: b for b in boxes}
    for b in boxes:
        col = 2 + 3 * b.level
        if b.kind == "orphan":
            side = Side(style="medium", color="C00000")
        elif b.status == "removed":
            side = Side(style="dashed", color="C00000")
        elif b.dashed:
            side = Side(style="dashed", color="808080")
        else:
            side = Side(style="thin", color="000000")
        key = b.status or ("orphan" if b.kind == "orphan" else ("dnp" if b.dashed else ""))
        fill = PatternFill("solid", fgColor=fills[key]) if key else None
        grey = b.dashed and b.status != "removed"
        for i, text in enumerate(b.lines):
            r = top0 + b.top + i
            c = ws.cell(r, col)
            link = pdf_link(b.page, text) if (pdf_link and i == 2 and b.page is not None and b.kind == "conv") else None
            c.value = link or text
            c.font = Font(name=FONT, size=9, bold=(i == 0), underline="single" if link else None,
                          color="0563C1" if link else ("808080" if grey else "000000"))
            c.alignment = Alignment(horizontal="left", vertical="center", shrink_to_fit=True)
            c.border = Border(left=side, right=side, top=side if i == 0 else None,
                              bottom=side if i == b.height - 1 else None)
            if fill:
                c.fill = fill
        if b.hover:
            cm = Comment("\n".join(b.hover), "AutoCheck")
            cm.width, cm.height = 420, 18 * (len(b.hover) + 1)
            ws.cell(top0 + b.top, col).comment = cm
    # connectors, once per parent: horizontal in the parent's first connector column, vertical trunk on the left
    # edge of the second connector column down to the last child, a stub into each child
    kids: dict[int, list[Box]] = {}
    for b in boxes:
        if b.parent is not None and b.level == by_id[b.parent].level + 1:
            kids.setdefault(b.parent, []).append(b)
    for pid, ch in kids.items():
        p = by_id[pid]
        pc = 2 + 3 * p.level
        prow = top0 + p.top + ANCHOR - 1              # bottom border of this row = the anchor line
        add_border(ws.cell(prow, pc + 1), bottom=line)
        last = prow
        for b in ch:
            crow = top0 + b.top + ANCHOR - 1
            add_border(ws.cell(crow, pc + 2), bottom=line)
            last = max(last, crow)
        for r in range(prow + 1, last + 1):
            add_border(ws.cell(r, pc + 2), left=line)
    if not boxes:
        ws.cell(top0, 2, "這份 netlist 沒有串出任何 power tree (沒有符合 pin_roles 的轉換元件)。").font = Font(name=FONT, size=10)
    ws.freeze_panes = ws.cell(top0, 1)


# --------------------------------------------------------------------------------------------------------- SVG
BW, LH, XGAP, PADX, PADY = 210, 16, 46, 16, 24
SVG_FILL = {"new": "#E2EFDA", "changed": "#FFF2CC", "removed": "#F8CBAD", "orphan": "#FDE9E7", "dnp": "#F2F2F2",
            "": "#FFFFFF"}


def render_svg(boxes: list[Box], pdf_url=None, a_label: str = "A", b_label: str = "B") -> str:
    """pdf_url(page) -> URL string or None. Same geometry as the Excel sheet (level, top, lines)."""
    by_id = {b.id: b for b in boxes}
    starts = section_starts(boxes)
    sec_extra = {sec: 18 * k for k, sec in enumerate(s for s in ("tree", "orphan", "removed") if s in starts)}

    def y_of(b: Box, line: int = 0) -> float:
        return PADY + sec_extra.get(b.section, 0) + (b.top + line) * LH

    def x_of(b: Box) -> float:
        return PADX + b.level * (BW + XGAP)

    maxlvl = max((b.level for b in boxes), default=0)
    w = PADX * 2 + (maxlvl + 1) * BW + maxlvl * XGAP
    h = max((y_of(b, b.height) for b in boxes), default=0) + PADY
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{w:.0f}" height="{h:.0f}" viewBox="0 0 {w:.0f} {h:.0f}" '
           'font-family="Arial, \'Microsoft JhengHei\', sans-serif" font-size="11">']
    for sec, r0 in starts.items():
        if sec in SECTION_TITLE:
            b0 = next(b for b in boxes if b.parent is None and b.section == sec and b.top == r0)
            out.append(f'<text x="{PADX}" y="{y_of(b0) - 8:.1f}" font-weight="bold" fill="#C00000">'
                       f'{escape(SECTION_TITLE[sec].format(a=a_label, b=b_label))}</text>')
    for b in boxes:                                   # connectors first, boxes on top
        if b.parent is None:
            continue
        p = by_id[b.parent]
        if b.level != p.level + 1:
            continue
        x0 = x_of(p) + BW
        x1 = x_of(b)
        xm = x0 + XGAP / 2
        y0 = y_of(p, ANCHOR)
        y1 = y_of(b, ANCHOR)
        dash = ' stroke-dasharray="4 3"' if b.dashed else ""
        col = "#C00000" if b.status == "removed" else "#404040"
        out.append(f'<path d="M{x0:.1f},{y0:.1f} H{xm:.1f} V{y1:.1f} H{x1:.1f}" fill="none" stroke="{col}" '
                   f'stroke-width="1.2"{dash}/>')
    for b in boxes:
        x, y = x_of(b), y_of(b)
        key = b.status or ("orphan" if b.kind == "orphan" else ("dnp" if b.dashed else ""))
        stroke, sw = "#000000", 1
        if b.kind == "orphan":
            stroke, sw = "#C00000", 2.2
        elif b.status == "removed":
            stroke = "#C00000"
        elif b.dashed:
            stroke = "#808080"
        dash = ' stroke-dasharray="5 3"' if (b.dashed or b.status == "removed") else ""
        url = pdf_url(b.page) if (pdf_url and b.page is not None and b.kind == "conv") else None
        g_open = f'<a href="{escape(url)}" target="_blank">' if url else ""
        g_close = "</a>" if url else ""
        out.append(f'<g class="box">{g_open}<title>{escape(chr(10).join(b.hover))}</title>'
                   f'<rect x="{x:.1f}" y="{y:.1f}" width="{BW}" height="{b.height * LH}" rx="3" fill="{SVG_FILL[key]}" '
                   f'stroke="{stroke}" stroke-width="{sw}"{dash}/>')
        grey = b.dashed and b.status != "removed"
        for i, t in enumerate(b.lines):
            attrs = ' font-weight="bold"' if i == 0 else ""
            fill = "#808080" if grey else "#000000"
            if i == 2 and url:
                fill, attrs = "#0563C1", attrs + ' text-decoration="underline"'
            out.append(f'<text x="{x + 6:.1f}" y="{y + (i + 1) * LH - 4:.1f}" fill="{fill}"{attrs}>'
                       f'{escape(_clip(t, 27 if i == 0 else 33))}</text>')
        out.append(f"{g_close}</g>")
    out.append("</svg>")
    return "\n".join(out)


def _clip(t: str, n: float = 33) -> str:
    """Clip to the box width: ASCII ~1 unit, CJK / full-width ~1.8 units at 11 px Arial (box text width 198 px)."""
    w, out = 0.0, []
    for ch in t:
        w += 1.8 if ord(ch) > 0x2E7F else 1.0
        if w > n - 1:
            return "".join(out) + "…"
        out.append(ch)
    return t


def render_html(boxes: list[Box], title: str, pdf_url=None, a_label: str = "A", b_label: str = "B",
                diff: bool = True) -> str:
    svg = render_svg(boxes, pdf_url, a_label, b_label)
    legend = "".join(
        f'<span class="lg"><i style="background:{SVG_FILL.get(k if k != "dnp" else "dnp", "#fff")};'
        f'border:{"2px solid #C00000" if k == "orphan" else ("1px dashed #C00000" if k == "removed" else ("1px dashed #808080" if k == "dnp" else "1px solid #000"))}"></i>{escape(t)}</span>'
        for k, t in LEGEND)
    note = "" if diff else f'<p class="warn">{escape(a_label)} 的 power tree 沒有算出來，本圖沒有變更顏色。</p>'
    return f"""<!doctype html>
<html lang="zh-Hant"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{escape(title)}</title>
<style>
:root {{ color-scheme: light; }}
body {{ margin: 0; padding: 12px 16px; background: #fff; color: #000; font: 13px Arial, 'Microsoft JhengHei', sans-serif; }}
h1 {{ font-size: 16px; margin: 0 0 6px; }}
p {{ margin: 4px 0; color: #444; }}
.warn {{ color: #C00000; }}
.lg {{ display: inline-flex; align-items: center; margin: 2px 14px 2px 0; font-size: 12px; }}
.lg i {{ display: inline-block; width: 18px; height: 12px; margin-right: 5px; }}
.wrap {{ overflow: auto; border-top: 1px solid #ddd; margin-top: 8px; padding-top: 8px; }}
.box:hover rect {{ stroke-width: 2.5; }}
</style></head><body>
<h1>{escape(title)}</h1>
<p>框內：rail 名 / 電壓 / 轉換元件 RefDes (藍字 = 點擊開線路圖該頁) / 負載 IC 數。滑鼠停在框上看料號、EN / PG、同組 net、負載清單。</p>
<div>{legend}</div>{note}
<div class="wrap">{svg}</div>
</body></html>
"""
