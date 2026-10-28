"""Excel tracker: one row per tracking row (see tracking.py), RD fills disposition columns.

Sheets: 版本資訊 (inputs, legend, PDF base URL cell), 追蹤表, 統計 (formulas), 規則說明, 電容耐壓表, 待討論.
The PDF column is =HYPERLINK(base_url & "#nameddest=" & dest): base_url lives in ONE cell on 版本資訊,
so moving the PDF (laptop -> web site) means editing one cell, not regenerating the file.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..rules_zh import DISPOSITIONS, LEVEL_ZH, RULES
from ..tracking import TrackRow

CAT_ZH = {"DRC": "基礎 DRC", "CAPV": "電容耐壓", "CONN": "連線變更", "PART": "零件變更",
          "BOM": "BOM 變更", "BOMCHK": "BOM 一致性", "NAMING": "命名規則", "NETRULE": "線路規則",
          "I2C": "I2C map"}
FONT = "Arial"
_thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=9)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
SEV_FILL = {"ERROR": PatternFill("solid", fgColor="F8CBAD"), "WARN": PatternFill("solid", fgColor="FFE699"),
            "INFO": PatternFill("solid", fgColor="E2EFDA")}

# 2026-10-27 B-full: 父項 column after 狀態 (= #n of the root-cause row this row is a consequence of; blank = top level).
# Every column position below is derived from HEADERS by NAME (COL): never hard-code a column letter or number here,
# review.fill_replies reads the sheet's own header row so older workbooks stay correct.
HEADERS = ["項次", "追蹤 Key", "群組 ID", "類別", "規則", "嚴重度", "狀態", "父項", "頁", "頁名", "RefDes", "Pin", "Net",
           "說明 (工具)", "細節 (工具)", "關聯 ID", "PDF", "PDF 頁", "處置", "負責人", "預計版次", "備註", "填寫日期",
           "PDF 目的地"]
WIDTHS = [6, 30, 14, 10, 20, 9, 10, 7, 5, 22, 10, 12, 26, 48, 60, 22, 6, 7, 10, 10, 10, 30, 11, 30]
assert len(WIDTHS) == len(HEADERS)
COL = {h: i + 1 for i, h in enumerate(HEADERS)}
INPUT_COLS = {COL[h] for h in ("處置", "負責人", "預計版次", "備註", "填寫日期")}
WRAP_COLS = {COL[h] for h in ("Net", "說明 (工具)", "細節 (工具)", "備註")}
PARENT_PREFIX = "同根因 → #{n}："      # rendered in front of a child row's 說明 (Excel + web); TrackRow.title itself is untouched


def _letter(h: str) -> str:
    return get_column_letter(COL[h])
BASE_URL_CELL = "$B$5"      # on 版本資訊


def _hdr(ws, row, headers, widths=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row, i, h)
        c.font = HDR_FONT; c.fill = HDR_FILL; c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]


def write_tracker(res, rows: list[TrackRow], out: Path, base_url: str = "", pdf_name: str = "",
                  sidecar: dict | None = None, questions: list[tuple[str, str]] | None = None,
                  replies: dict[str, dict] | None = None) -> Path:
    """replies: tracking key -> {處置, 負責人, 預計版次, 備註, 填寫日期} already given on the web page."""
    meta_a, meta_b = res.a, res.b
    pdf_pages = {k: v.get("pdf_page") for k, v in (sidecar or {}).get("rows", {}).items()}
    box_bad = {k for k, v in (sidecar or {}).get("rows", {}).items() if v.get("box_check") == "mismatch"}
    wb = Workbook()

    # ---- 版本資訊 ----
    ws = wb.active; ws.title = "版本資訊"
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 100

    def kv(r, k, v, fill=None):
        ws.cell(r, 1, k).font = Font(name=FONT, bold=True, size=10)
        c = ws.cell(r, 2, v); c.font = Font(name=FONT, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            c.fill = fill
    ws.cell(1, 1, f"AutoCheck 追蹤表  {meta_a.label} -> {meta_b.label}").font = Font(name=FONT, bold=True, size=14)
    kv(2, "產生時間", res.generated)
    kv(3, "工具 / 規則版本", f"tool {res.tool_version} | rules {res.cfg.get('rules_version')} | config {res.cfg.get('_hash')}")
    kv(4, "標註 PDF 檔名", pdf_name or "(本次未產生標註 PDF)")
    kv(5, "PDF base URL (可改)", base_url, INPUT_FILL)
    ws.cell(5, 2).comment = Comment("追蹤表 PDF 欄的連結 = 這格 + #nameddest=目的地。搬到網站後只改這格。", "AutoCheck")
    r = 7
    for d in (meta_a, meta_b):
        for s in d.sources:
            kv(r, f"{d.label} {s.role}", f"{Path(s.path).name}  |  SHA-256 {s.sha256[:12]}  |  {s.stamp}"); r += 1
    g0 = getattr(res, "bom_check", None) or {}
    for side, d in (("A", meta_a), ("B", meta_b)):      # G0 BOM / netlist consistency, both sides (owner 2026-10-22)
        line = _g0_line((g0.get(side) or {}).get("stats"), d.bom is not None)
        if line:
            bad = ((g0.get(side) or {}).get("stats") or {}).get("strong")
            kv(r, f"{d.label} BOM vs netlist", line, PatternFill("solid", fgColor="FCE4D6") if bad else None); r += 1
    r += 1
    ws.cell(r, 1, "填寫說明").font = Font(name=FONT, bold=True, size=12); r += 1
    for k, v in [
        ("黃底欄位", "RD 填寫：處置、負責人、預計版次、備註、填寫日期。其他欄位由工具產生，不要改。"),
        ("處置選項", "修改 = 真 bug，下版改；刻意設計 = 現況正確，下版重跑自動標「已確認」；誤報 = 工具規則要調；待討論；N/A = 不適用。"),
        ("追蹤 Key", "工具回讀 waiver 的唯一鍵 = 規則 | RefDes (或 pin / 具名 net)。不含自動命名 net 與群組 ID，改版重匯 netlist 後仍穩定。"),
        ("項次 #", "與標註 PDF 上框框旁的 #n、索引頁的 # 相同。"),
        ("群組 ID", "PDF 報告中的 finding ID，同一根因的多顆零件共用；成員一變 ID 就變，只拿來對照報告，不能當追蹤鍵。"),
        ("PDF 欄", "點「PDF」跳到標註 PDF 的該筆位置 (需 base URL 正確且瀏覽器開啟)。「PDF 頁」是標註 PDF 的頁碼 (含索引頁)。"),
        ("狀態", "NEW 本版新增 / EXISTING 舊版已存在 / CHANGE 改版差異 (A->B)。"),
        ("父項", "同根因的根源列 #n (例：短路 NET_SHORTED 底下的 NET_SPLIT / PIN_MOVED / 單腳 net)。子項說明前加「同根因 → #n」，"
                "PDF 上不另外框 (嚴重度高於父項者與 PIN_ISOLATED 例外)；處置各列獨立填，統計表進度不含子項。篩「父項 = 空白」只看根源列。"),
        ("嚴重度", "ERROR 疑似缺陷 / WARN 需審查 / INFO 備查 (預設可不處置)。"),
    ]:
        kv(r, k, v); r += 1

    # ---- 追蹤表 ----
    wt = wb.create_sheet("追蹤表")
    _hdr(wt, 1, HEADERS, WIDTHS)
    wt.row_dimensions[1].height = 30
    wt.cell(1, COL["追蹤 Key"]).comment = Comment("waiver 回讀用的穩定鍵。", "AutoCheck")
    wt.cell(1, COL["處置"]).comment = Comment("下拉選單：" + " / ".join(DISPOSITIONS), "AutoCheck")
    rowno = {r.key: i for i, r in enumerate(rows, 1)}
    wt.cell(1, COL["父項"]).comment = Comment("同根因的根源列 #n；空白 = 最上層。子項不在 PDF 上另外框 (嚴重度高於父項者除外)。", "AutoCheck")
    for i, row in enumerate(rows, 2):
        n = i - 1
        pn = rowno.get(row.parent) if getattr(row, "parent", "") else None
        title = (PARENT_PREFIX.format(n=pn) + row.title) if pn else row.title
        vals = [n, row.key, row.gid, row.category, row.rule, row.severity, row.status, f"#{pn}" if pn else "",
                row.page, row.page_title, ", ".join(row.refdes), ", ".join(row.pins), ", ".join(row.nets), title,
                row.detail, ", ".join(row.related), None, pdf_pages.get(row.key), "", "", "", "", "", row.dest]
        assert len(vals) == len(HEADERS)
        for j, v in enumerate(vals, 1):
            c = wt.cell(i, j, v); c.font = BODY; c.border = BORDER
            c.alignment = Alignment(wrap_text=(j in WRAP_COLS), vertical="top")
            if j in INPUT_COLS:
                c.fill = INPUT_FILL
        wt.cell(i, COL["嚴重度"]).fill = SEV_FILL.get(row.severity, SEV_FILL["INFO"])
        d = get_column_letter(COL["PDF 目的地"])
        c = wt.cell(i, COL["PDF"], f'=IF(OR(版本資訊!{BASE_URL_CELL}="",{d}{i}=""),"",'
                                  f'HYPERLINK(版本資訊!{BASE_URL_CELL}&"#nameddest="&{d}{i},"PDF"))')
        c.font = Font(name=FONT, size=9, color="0563C1", underline="single")
        if row.key in box_bad:          # annotate self-check failed: mark, never block
            pc = wt.cell(i, COL["PDF 頁"])
            pc.fill = PatternFill("solid", fgColor="FFC7CE")
            pc.comment = Comment("框位置自檢未通過：標註 PDF 上的框可能沒落在零件旁，請以頁碼人工找零件。", "AutoCheck")
    last = len(rows) + 1
    dv = DataValidation(type="list", formula1='"' + ",".join(DISPOSITIONS) + '"', allow_blank=True)
    wt.add_data_validation(dv)
    dv.add(f"{get_column_letter(COL['處置'])}2:{get_column_letter(COL['處置'])}{max(last, 2000)}")
    wt.freeze_panes = f"{_letter('Pin')}2"          # frozen through RefDes: Pin is the first scrolling column (was "K2")
    wt.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(last, 2)}"
    wt.column_dimensions[get_column_letter(COL["PDF 目的地"])].hidden = True
    if replies:
        from ..review import fill_replies
        fill_replies(wt, replies)

    # ---- 統計 ----
    wsum = wb.create_sheet("統計")
    wsum.column_dimensions["A"].width = 22
    for col in "BCDEFG":
        wsum.column_dimensions[col].width = 14
    wsum.cell(1, 1, "發現數 (追蹤表列數，群組已展開)").font = Font(name=FONT, bold=True, size=12)
    _hdr(wsum, 2, ["類別", "ERROR", "WARN", "INFO", "小計", "NEW (ERROR+WARN)"])
    lo, hi = 2, max(last, 2)
    def rng(h: str) -> str:
        return f"追蹤表!${_letter(h)}${lo}:${_letter(h)}${hi}"
    T, S, ST, P, PA = rng("類別"), rng("嚴重度"), rng("狀態"), rng("處置"), rng("父項")
    r = 3
    for cat, zh in CAT_ZH.items():
        wsum.cell(r, 1, f"{cat} {zh}")
        for j, sev in enumerate(("ERROR", "WARN", "INFO"), 2):
            wsum.cell(r, j, f'=COUNTIFS({T},"{cat}",{S},"{sev}")')
        wsum.cell(r, 5, f"=SUM(B{r}:D{r})")
        wsum.cell(r, 6, f'=COUNTIFS({T},"{cat}",{ST},"NEW",{S},"ERROR")+COUNTIFS({T},"{cat}",{ST},"NEW",{S},"WARN")')
        for j in range(1, 7):
            wsum.cell(r, j).border = BORDER; wsum.cell(r, j).font = BODY
        r += 1
    wsum.cell(r, 1, "合計").font = Font(name=FONT, bold=True, size=9); wsum.cell(r, 1).border = BORDER
    for j in range(2, 7):
        col = get_column_letter(j)
        c = wsum.cell(r, j, f"=SUM({col}3:{col}{r - 1})"); c.font = Font(name=FONT, bold=True, size=9); c.border = BORDER
    r += 2
    wsum.cell(r, 1, "處置進度 (不含子項：父項欄非空白的列，同根因只算根源那一列)").font = Font(name=FONT, bold=True, size=12); r += 1
    _hdr(wsum, r, ["處置", "列數", "其中 ERROR", "其中 WARN"]); r += 1
    for disp in DISPOSITIONS + ["(未填)"]:
        q = '""' if disp == "(未填)" else f'"{disp}"'
        wsum.cell(r, 1, disp)
        wsum.cell(r, 2, f'=COUNTIFS({P},{q},{PA},"")')
        wsum.cell(r, 3, f'=COUNTIFS({S},"ERROR",{P},{q},{PA},"")')
        wsum.cell(r, 4, f'=COUNTIFS({S},"WARN",{P},{q},{PA},"")')
        for j in range(1, 5):
            wsum.cell(r, j).border = BORDER; wsum.cell(r, j).font = BODY
        r += 1
    wsum.cell(r, 1, "總列數").font = BODY; wsum.cell(r, 2, f"=COUNTA(追蹤表!$A${lo}:$A${hi})").font = BODY
    r += 1
    wsum.cell(r, 1, "其中子項列數").font = BODY; wsum.cell(r, 2, f'=COUNTIF({PA},"#*")').font = BODY
    cov = (getattr(res, "coverage", {}) or {}).get("DECOUPLE")
    if cov:   # C1 coverage (QA 2026-10-20): what the rule looked at, so 0 finding is not mistaken for 100 % checked
        r += 2
        wsum.cell(r, 1, f"規則覆蓋率 DECOUPLE ({meta_b.label})").font = Font(name=FONT, bold=True, size=12); r += 1
        _hdr(wsum, r, ["項目", "數量", "說明"]); r += 1
        for lab, key, note in (("檢查的電源 net 數", "supply_nets", "餵上件 IC 的電源 net (排除 ground / NC / 自動命名)"),
                               ("認到的 IC 電源腳數", "supply_pins", "pin 角色在 config netrules.decouple_roles 的上件 IC 腳"),
                               ("電源腳接自動命名 net (未檢查)", "supply_pins_unnamed_net", "net 沒取名 (Nxxxxx)，C1 跳過；見下列"),
                               ("上件 IC 數", "ics_built", "BOM 有的 U*"),
                               ("沒認到任何電源腳的上件 IC 數", "ics_no_supply_pin", "這些 IC 的 #14 沒檢查到，見下列")):
            wsum.cell(r, 1, lab); wsum.cell(r, 2, cov.get(key, 0)); wsum.cell(r, 3, note)
            for j in range(1, 4):
                wsum.cell(r, j).border = BORDER; wsum.cell(r, j).font = BODY
            r += 1
        for lab, key in (("未覆蓋 IC", "ics_no_supply_pin_list"), ("電源腳在自動命名 net", "supply_pins_unnamed_list")):
            lst = cov.get(key) or []
            if not lst:
                continue
            wsum.cell(r, 1, lab).font = BODY
            # 20 per row: one cell holds at most 32767 chars (QA #13)
            for i in range(0, len(lst), 20):
                wsum.cell(r, 2, ", ".join(lst[i:i + 20])).font = BODY
                wsum.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
                wsum.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
                r += 1
    cov = (getattr(res, "coverage", {}) or {}).get("FET")
    if cov:   # C3 FET_GATE_LEVEL coverage (QA 2026-10-21 L2): which FETs got a row, which were skipped
        r += 2
        wsum.cell(r, 1, f"規則覆蓋率 FET_GATE_LEVEL ({meta_b.label})").font = Font(name=FONT, bold=True, size=12); r += 1
        _hdr(wsum, r, ["項目", "數量", "說明"]); r += 1
        for lab, key, note in (("上件 Q* 零件數", "fets_built", "RefDes 前綴在 config netrules.fet_prefixes 的上件零件 (含 BJT)"),
                               ("列出的電晶體數", "transistors_listed", "每顆 FET (dual 每半顆) 一列 Vg / Vs / Vgs"),
                               ("沒拆出 gate / source 的零件數", "fets_skipped", "BJT、無後綴的 dual (G/S/D x2)、腳名不在 pin_roles：未列，見下列")):
            wsum.cell(r, 1, lab); wsum.cell(r, 2, cov.get(key, 0)); wsum.cell(r, 3, note)
            for j in range(1, 4):
                wsum.cell(r, j).border = BORDER; wsum.cell(r, j).font = BODY
            r += 1
        lst = cov.get("fets_skipped_list") or []
        if lst:
            wsum.cell(r, 1, "未列的零件").font = BODY
            for i in range(0, len(lst), 20):
                wsum.cell(r, 2, ", ".join(lst[i:i + 20])).font = BODY
                wsum.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
                wsum.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
                r += 1

    # ---- 規則說明 ----
    wr = wb.create_sheet("規則說明")
    _hdr(wr, 1, ["規則", "類別", "追蹤粒度", "中文名稱", "為什麼 / 怎麼查", "本次筆數"], [24, 10, 10, 22, 90, 10])
    for i, (rule, (cat, level, zh, why)) in enumerate(RULES.items(), 2):
        for j, v in enumerate([rule, cat, LEVEL_ZH[level], zh, why, f'=COUNTIF({rng("規則")},A{i})'], 1):
            c = wr.cell(i, j, v); c.font = BODY; c.border = BORDER; c.alignment = Alignment(wrap_text=True, vertical="top")

    # ---- 電容耐壓表 ----
    if res.cap_rows:
        wc = wb.create_sheet(f"電容耐壓表 ({meta_b.label})")
        fields = list(res.cap_rows[0])
        _hdr(wc, 1, fields, [13] * len(fields))
        for ri, row in enumerate(res.cap_rows, 2):
            for ci, k in enumerate(fields, 1):
                c = wc.cell(ri, ci, row.get(k)); c.font = BODY; c.border = BORDER
        wc.freeze_panes = "B2"; wc.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{len(res.cap_rows) + 1}"

    # ---- BOM 一致性 (G0, 2026-10-22): every part whose BOM line and netlist disagree, A and B ----
    _g0_sheet(wb, [(meta_a.label, (g0.get("A") or {})), (meta_b.label, (g0.get("B") or {}))])

    # ---- Power tree (B1) ----
    pt = getattr(res, "power_tree", None) or {}
    if pt.get("rows"):
        wp = wb.create_sheet(f"Power tree ({meta_b.label})")
        cols = [("level", "層級", 6), ("tree", "樹狀 (下游 rail)", 30), ("rail_in", "上游 rail", 22), ("v_in", "上游 V", 8),
                ("conv", "轉換元件", 9), ("mpn", "料號", 20), ("kind", "類型", 18), ("built", "上件", 6),
                ("v_out", "下游 V", 8), ("v_src", "V 來源", 9), ("en", "EN net", 26), ("pg", "PG net", 22), ("via", "經過 (電感 / FET)", 26),
                ("page", "頁", 5), ("out_members", "下游 rail 同組 net (經 0R / fuse / 串阻)", 60)]
        wp.cell(1, 1, "從 netlist pin 名稱自動串出的 power tree (B1)。請拿線路圖裡的 power tree 圖對照；"
                      "工具認不出來源的 rail 列在最下面。").font = Font(name=FONT, bold=True, size=10)
        _hdr(wp, 2, [c[1] for c in cols], [c[2] for c in cols])
        r = 3
        for row in pt["rows"]:
            for j, (k, _, _) in enumerate(cols, 1):
                c = wp.cell(r, j, row.get(k)); c.font = BODY; c.border = BORDER
                c.alignment = Alignment(vertical="top", wrap_text=k in ("out_members", "via", "en"))
                if row.get("built") == "DNP":
                    c.font = Font(name=FONT, size=9, color="808080")
            r += 1
        if pt.get("orphans"):
            r += 1
            wp.cell(r, 1, "有負載但工具認不出來源的 rail (含 converter 內部 net：工具判斷欄寫 internal net 的，方框圖預設不畫)").font = Font(name=FONT, bold=True, size=10); r += 1
            ocols = [("rail", "rail"), ("voltage", "V"), ("v_src", "V 來源"), ("loads", "負載 IC (電源腳)"), ("source", "工具判斷"),
                     ("members", "同組 net")]
            _hdr(wp, r, [c[1] for c in ocols]); r += 1
            for o in pt["orphans"]:
                for j, (k, _) in enumerate(ocols, 1):
                    c = wp.cell(r, j, o.get(k)); c.font = BODY; c.border = BORDER
                    c.alignment = Alignment(vertical="top", wrap_text=True)
                r += 1
        wp.freeze_panes = "C3"

        # ---- Power tree 圖: the same data as boxes (owner request 2026-10-12) ----
        if pt.get("edges") is not None:
            from .power_tree_diagram import layout, render_xlsx
            pt_a = getattr(res, "power_tree_a", None) or None
            diff = "edges" in (pt_a or {})
            # converter-internal nets hidden (owner 2026-10-13): listed in the converter box comment and in the
            # "Power tree" table sheet; the web page has the toggle
            boxes = layout(pt, pt_a, meta_a.label, meta_b.label, show_internal=False)
            pmap = (sidecar or {}).get("page_map", {})
            link = None
            if pmap:
                def link(page, text):
                    n = pmap.get(str(page))
                    if not n:
                        return None
                    t = text.replace('"', '""')
                    # same guard as the 追蹤表 PDF column: empty base URL cell -> plain text, no dead link
                    return (f'=IF(版本資訊!{BASE_URL_CELL}="","{t}",'
                            f'HYPERLINK(版本資訊!{BASE_URL_CELL}&"#page={n}","{t}"))')
            wg = wb.create_sheet(f"Power tree 圖 ({meta_b.label})"[:31])     # Excel sheet names: max 31 chars
            render_xlsx(wg, boxes, f"Power tree 方框圖 ({meta_b.label})" + (
                f"，顏色 = {meta_a.label} -> {meta_b.label} 變更" if diff else ""), link, meta_a.label, meta_b.label,
                diff=diff, banner=(pt.get("bom_check") or {}).get("banner", ""))

    # ---- I2C map (2026-10-12, test state) ----
    im = getattr(res, "i2c_map", None) or {}
    if im.get("buses"):
        from ..checks.i2c_map import table_rows
        wi = wb.create_sheet(f"I2C map ({meta_b.label})"[:31])
        wi.cell(1, 1, "從 netlist + BOM + config/i2c_devices.yaml 推出的 I2C map (測試狀態，規則全部 INFO)。strap 狀態："
                      "0 = GND、1 = 接到 rail、Z = 空接、SDA / SCL = 接自己的 SDA / SCL、X = 上下拉都上件、? = rail 電壓不明；"
                      "→ 後面是算出的 7-bit 位址，strap 電阻欄列每支位址腳實際上件 / DNP 的電阻。"
                      "方框圖：輸出資料夾 i2c_map_<版>.html / UI「I2C map」頁。").font = Font(name=FONT, bold=True, size=10)
        cols = [("bus", "bus", 18), ("speed", "速度 (標示)", 14), ("refdes", "RefDes", 8), ("symbol", "symbol 料號", 20),
                ("bom", "BOM 主料", 30), ("built", "上件", 6), ("addr7", "7-bit", 9), ("addr8", "8-bit W / R", 12),
                ("addr_line", "strap 狀態 → 位址", 30), ("strap", "strap 電阻 (上件)", 40),
                ("strap_dnp", "strap 電阻 (DNP)", 36), ("note", "位址備註", 16), ("lib", "I2C 庫", 24), ("series", "分支串接件", 26),
                ("sda", "SDA pin", 20), ("scl", "SCL pin", 16), ("scope", "位址空間", 16), ("page", "頁", 5),
                ("change", "變更", 40)]
        _hdr(wi, 2, [c[1] for c in cols], [c[2] for c in cols])
        r = 3
        for row in table_rows(im, meta_a.label):
            for j, (k, _, _) in enumerate(cols, 1):
                v = row.get(k)
                if k in ("strap", "strap_dnp") and v:
                    v = v.replace("; ", "\n")               # one address pin per line in the cell
                c = wi.cell(r, j, v); c.font = BODY; c.border = BORDER
                c.alignment = Alignment(vertical="top", wrap_text=k in ("addr_line", "strap", "strap_dnp", "series", "sda", "bom", "change"))
                if row.get("built") == "DNP":
                    c.font = Font(name=FONT, size=9, color="808080")
            r += 1
        if im.get("removed"):
            r += 1
            wi.cell(r, 1, f"{meta_a.label} 有、本版沒有的 I2C 裝置: {', '.join(im['removed'])}").font = Font(
                name=FONT, size=9, color="C00000")
        # pull-ups: own sheet (column widths of the device table do not fit it)
        wu = wb.create_sheet(f"I2C pull-up ({meta_b.label})"[:31])
        wu.cell(1, 1, "I2C pull-up，每條 segment (上件的 0R / 串阻連在一起的 nets) 每條線。同一條線上件超過一顆 = 紅底 "
                      "(I2C_PU_DUPLICATE)。DNP 不算。").font = Font(name=FONT, bold=True, size=10)
        _hdr(wu, 2, ["bus", "segment", "線", "上件 pull-up", "DNP pull-up", "連接器 (出板，net name 為主)", "裝置", "nets"],
             [20, 24, 6, 34, 34, 34, 30, 60])

        def _conn_txt(b, sg) -> str:
            # owner 2026-10-12: connector judged by net name; the symbol pin name (PETN0 on a gold finger) is a note only
            out = [f"{net} ({pin.split(' ', 1)[0]}" + (f"，pin 名 {pin.split(' ', 1)[1]} 僅附註)" if " " in pin else ")")
                   for ln, net, pin in getattr(b, "conn_pins", []) if net in sg.nets]
            return "\n".join(out) if out else ", ".join(sg.connectors)
        r = 3
        for b in im["buses"]:
            for sg in b.segments:
                vals = [b.bid, sg.sid.split(":", 1)[1], sg.line, "; ".join(sg.built_pullups) or "(無)",
                        "; ".join(sg.dnp_pullups), _conn_txt(b, sg), ", ".join(sg.devices), ", ".join(sg.nets)]
                for j, v in enumerate(vals, 1):
                    c = wu.cell(r, j, v); c.font = BODY; c.border = BORDER
                    c.alignment = Alignment(vertical="top", wrap_text=True)
                if len(sg.built_pullups) > 1:
                    wu.cell(r, 4).fill = PatternFill("solid", fgColor="FFC7CE")
                r += 1
        wu.freeze_panes = "C3"
        # I2C Block Diagram comparison: own sheet
        decl = im.get("declared") or []
        wd = wb.create_sheet(f"I2C 圖對照 ({meta_b.label})"[:31])
        if decl:
            n_bad = sum(1 for x in decl if x["verdict"] == "不符")
            wd.cell(1, 1, f"線路圖 I2C Block Diagram (轉錄檔 {Path(im.get('declared_file', '')).name}) 逐列對照 netlist / BOM："
                          f"{len(decl)} 列，不符 {n_bad} 列 (紅)，無法驗證 = 黃").font = Font(name=FONT, bold=True, size=10)
            dcols = [("no", "#", 4), ("kind", "類型", 9), ("label", "圖上名稱", 16), ("refdes", "圖上 RefDes", 9),
                     ("mpn", "圖上料號", 16), ("addr7", "圖上 7-bit", 8), ("net", "net", 34), ("value", "值", 7),
                     ("dnp", "圖上 DNP", 7), ("actual", "netlist 實際", 12), ("verdict", "結果", 9), ("why", "說明", 90)]
            _hdr(wd, 2, [c[1] for c in dcols], [c[2] for c in dcols])
            r = 3
            for row in decl:
                for j, (k, _, _) in enumerate(dcols, 1):
                    c = wd.cell(r, j, row.get(k)); c.font = BODY; c.border = BORDER
                    c.alignment = Alignment(vertical="top", wrap_text=k in ("why", "net"))
                if row["verdict"] == "不符":
                    wd.cell(r, 11).fill = PatternFill("solid", fgColor="FFC7CE")
                elif row["verdict"] == "無法驗證":
                    wd.cell(r, 11).fill = PatternFill("solid", fgColor="FFEB9C")
                r += 1
            wd.freeze_panes = "D3"
        else:
            wd.cell(1, 1, "沒有 I2C Block Diagram 對照檔：把簽核過的 i2c_map_declared*.csv 放進新版資料夾 (或 --i2c-declared / UI 欄位)。"
                    ).font = Font(name=FONT, size=10, color="C00000")
        wi.freeze_panes = "D3"

    # ---- 待討論 ----
    if questions:
        wq = wb.create_sheet("待討論")
        _hdr(wq, 1, ["#", "議題", "草稿立場 / 選項", "結論"], [5, 40, 90, 30])
        for i, (t, s) in enumerate(questions, 2):
            for j, v in enumerate([i - 1, t, s, ""], 1):
                c = wq.cell(i, j, v); c.font = BODY; c.border = BORDER; c.alignment = Alignment(wrap_text=True, vertical="top")
            wq.cell(i, 4).fill = INPUT_FILL

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def _g0_line(st: dict | None, has_bom: bool) -> str:
    """One 版本資訊 line of G0 statistics ("" when the side has no BOM)."""
    if not st or not has_bom:
        return ""
    c = st.get("status") or {}
    px = ", ".join(f"{k} {v}" for k, v in (st.get("by_prefix") or {}).items())
    s = (f"比對 {st.get('compared', 0)} 顆 (BOM 與 netlist 都有)：料號相同 {c.get('OK_PN', 0)}、值相同 {c.get('OK_VALUE', 0)}、"
         f"同值換料號 {c.get('SAME_VALUE', 0)}、不符 {st.get('mismatch', 0)} (值不同 {c.get('VALUE_DIFF', 0)} / "
         f"類型不同 {c.get('TYPE_DIFF', 0)} / 料號不同 {c.get('PN_DIFF', 0)})、無法比對 {c.get('UNCHECKED', 0)}")
    if px:
        s += f"；不符依前綴 {px}"
    strong = st.get("strong") or []
    if strong:
        s += (f"；{'/'.join(st.get('strong_prefixes') or [])} 不符 {len(strong)} 顆 = power tree ⚠"
              f" ({', '.join(strong[:12])}{' …' if len(strong) > 12 else ''})")
    s += f"；netlist 有 BOM 無 (DNP) {st.get('netlist_not_in_bom', 0)}、BOM 有 netlist 無 {st.get('bom_not_in_netlist', 0)}"
    return s


def _g0_sheet(wb, sides: list[tuple[str, dict]]) -> None:
    from ..checks.bom_consistency import STATUS_ZH, sheet_rows
    rows = sheet_rows(sides)
    if not any((chk or {}).get("rows") for _, chk in sides):
        return
    ws = wb.create_sheet("BOM 一致性")
    ws.cell(1, 1, "BOM 與 netlist 逐顆比對 (G0)。只列不是「料號相同 / 值相同」的料；同值換料號只是統計，不進追蹤表。"
                  "比法：有 NV_PN 等料號屬性比 BOM 料號 + AML 替代料，沒有就比符號值 / MPN 文字。").font = Font(name=FONT, size=9)
    heads = ["板", "RefDes", "狀態", "差異", "比對方式", "netlist 料號", "netlist 值", "netlist 換算", "BOM 料號",
             "BOM 描述", "BOM 規格 / 廠牌", "BOM 換算", "AML 替代料", "BOM 列", "頁"]
    _hdr(ws, 2, heads, [8, 14, 12, 18, 9, 16, 22, 12, 16, 40, 28, 12, 22, 7, 5])
    for i, x in enumerate(rows, 3):
        vals = [x["side"], x["refdes"], STATUS_ZH.get(x["status"], x["status"]), x.get("what", ""),
                "料號" if x["mode"] == "pn" else "文字 / 值", x.get("netlist_pn", ""), x.get("netlist_value", ""),
                x.get("netlist_fmt", ""), x.get("bom_pn", ""), x.get("bom_desc", ""), x.get("bom_spec", ""),
                x.get("bom_fmt", ""), ", ".join(x.get("alternates") or []), x.get("bom_row"), x.get("page")]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v); c.font = BODY; c.border = BORDER
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(heads))}{max(len(rows) + 2, 3)}"
