"""Browser adapter for the user's original AutoCheck 2026-10-27 engine."""
import dataclasses
import datetime
import json
import os
import re
from html import escape
from pathlib import Path

import yaml

from autocheck import review
from autocheck.checks import i2c_map as i2c
from autocheck.cli import write_outputs
from autocheck.config import DEFAULT, load_config
from autocheck.i2c_lib import accept_entry, draft_from_pdf, entry_yaml, page_text, validate_entry
from autocheck.pipeline import run
from autocheck.report import pdf_report
from autocheck.report.data_out import meta
from autocheck.tracking import TrackRow, expand_rows, link_root_causes
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont


pdfmetrics.registerFont(UnicodeCIDFont('MSung-Light'))
pdf_report._FONT.update(name='MSung-Light', bold='MSung-Light', cjk=True)


def _folder_datasheets(folder, texts):
    """Keep browser-selected subfolders and recognize PDF extension casing."""
    if not folder.is_dir():
        return []
    norm = [i2c._norm(text) for text in texts]
    paths = sorted(
        (path for path in folder.rglob('*') if path.is_file() and path.suffix.lower() == '.pdf'),
        key=lambda path: (path.relative_to(folder).as_posix().casefold(), path.relative_to(folder).as_posix()),
    )
    return [path for path in paths if len(stem := i2c._norm(path.stem)) >= 5
            and any(stem in text for text in norm)]


# The original engine's hint resolves this module attribute at call time. Extend
# only browser file discovery; the engine source and its integrity manifest stay
# unchanged, and no uploaded PDF becomes a signed library rule automatically.
i2c.matching_datasheets = _folder_datasheets


def _plain(value):
    if dataclasses.is_dataclass(value):
        return _plain(dataclasses.asdict(value))
    if isinstance(value, dict):
        return {str(k): _plain(v) for k, v in value.items() if k != '_findings'}
    if isinstance(value, (list, tuple, set)):
        return [_plain(v) for v in value]
    if isinstance(value, re.Pattern):
        return value.pattern
    if isinstance(value, (Path, datetime.date)):
        return str(value)
    return value


def _label(value, default):
    value = str(value or default).strip()
    if len(value) > 80 or any(c in value for c in '/\\[]:*?\x00') or any(ord(c) < 32 for c in value):
        raise ValueError('版次名稱不可包含 / \\ [ ] : * ? 或控制字元，且不得超過 80 字')
    return value


def _yaml_doc(text):
    if not isinstance(text, str) or len(text.encode('utf-8')) > 2 * 1024 * 1024:
        raise ValueError('YAML 必須是文字且不得超過 2 MiB')
    doc = yaml.safe_load(text) or {}
    if not isinstance(doc, dict):
        raise ValueError('YAML 最上層必須是欄位對照')
    count = 0
    active = set()

    def check(value, depth=0):
        nonlocal count
        count += 1
        if count > 100000 or depth > 64:
            raise ValueError('YAML 結構過大或過深')
        if isinstance(value, (list, dict)):
            if id(value) in active:
                raise ValueError('YAML 不可包含循環別名')
            active.add(id(value))
            for child in value.values() if isinstance(value, dict) else value:
                check(child, depth + 1)
            active.remove(id(value))
    check(doc)
    return doc


def _option_text(value):
    return Path(value['file']).read_text(encoding='utf-8-sig') if isinstance(value, dict) else value


def _library(payload, root):
    path = root / 'i2c_devices_user.yaml'
    text = payload.get('yaml')
    text = 'devices: []\n' if text is None else text
    errors, warnings = [], []
    try:
        doc = _yaml_doc(text)
        items = doc.get('devices')
        if not isinstance(items, list):
            raise ValueError('使用者位址庫必須包含 devices 清單')
        for n, item in enumerate(items, 1):
            if not isinstance(item, dict):
                errors.append(f'第 {n} 筆位址規則格式無效')
                continue
            err, warn = validate_entry(item)
            errors.extend(f'第 {n} 筆：{value}' for value in err)
            warnings.extend(f'第 {n} 筆：{value}' for value in warn)
        if not errors:
            path.write_text(text, encoding='utf-8')
            candidate = _yaml_doc(payload['entryYaml']) if payload.get('entryYaml') is not None else payload.get('entry')
            if candidate is not None:
                if not isinstance(candidate, dict):
                    raise ValueError('候選規則格式無效')
                _, warn = validate_entry(candidate, items)
                warnings.extend(warn)
                accept_entry(candidate, payload.get('signedBy'), path)
                text = path.read_text(encoding='utf-8')
    except (ValueError, TypeError, yaml.YAMLError) as error:
        errors.append(str(error))
    cfg = load_config()
    cfg.setdefault('i2c', {})['user_library'] = str(path)
    entries = [dict(item.raw, origin=item.origin, verified=item.verified) for item in i2c.load_library(cfg)]
    warnings.extend(i2c.LIB_WARNINGS)
    return {'entries': entries, 'yaml': text, 'errors': errors, 'warnings': warnings}


_ROLE_SUFFIX = {'.xls': 'bom', '.xlsx': 'bom', '.htm': 'bom', '.html': 'bom', '.csv': 'bom', '.dsn': 'dsn', '.pdf': 'pdf'}


def _role(path):
    """The website already lets the user pick one file per role and sends only those, so roles follow the
    file names (the 2026-10-15 loader rules). The 2026-10-22 engine's loader chooses by content across a
    whole design-kit folder and needs all three netlist files together, which single-file calls never satisfy."""
    name = path.name.lower()
    if name.startswith('i2c_map_declared') and path.suffix.lower() == '.csv':
        return 'i2c_declared'
    for key in ('pstxnet', 'pstxprt', 'pstchip'):
        if key in name and name.endswith('.dat'):
            return key
    if path.suffix.lower() == '.pdf' and name.startswith('autocheck'):
        return None
    return _ROLE_SUFFIX.get(path.suffix.lower())


def _classified(paths, label):
    seen = {}
    for value in paths:
        path = Path(value)
        role = _role(path)
        if role is None:
            continue
        if role in seen:
            raise ValueError(f'{label} 有多個 {role} 候選，請先選擇要使用的檔案')
        seen[role] = path
    missing = [name for name in ('pstxnet', 'pstxprt', 'pstchip') if name not in seen]
    if missing:
        raise ValueError(f'{label} 缺少必要 netlist：{", ".join(missing)}')
    return seen


def _artifact(path, key=None):
    mime = {'.pdf': 'application/pdf', '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            '.json': 'application/json', '.csv': 'text/csv', '.txt': 'text/plain', '.html': 'text/html',
            '.drawio': 'application/vnd.jgraph.mxfile',
            }.get(path.suffix, 'application/octet-stream')
    return {'key': key or path.stem, 'name': path.name, 'mime': mime, 'path': str(path)}


def _drawio_outputs(result, out, paths):
    """Keep the upstream B renderers and expose independent A/B editable copies.

    The CLI owns B report generation, including its isolated renderer failures.
    A uses the same shared layouts without a comparison overlay. Browser-only
    key names never change the original engine or replace existing diagrams.
    """
    for original, key in (('ptree_drawio', 'power_drawio_b'), ('i2c_block_drawio', 'i2c_drawio_b')):
        path = paths.pop(original, None)
        if path is not None and path.is_file() and path.stat().st_size:
            paths[key] = path

    def write_a(kind, filename, make):
        # A and B labels may be equal or contain suffixes; never overwrite B's
        # existing file even when its label happens to match the A filename.
        target = out / filename
        suffix = 2
        while target.exists():
            target = out / f'{Path(filename).stem}_{suffix}.drawio'
            suffix += 1
        try:
            xml = make()
            target.write_text(xml, encoding='utf-8')
            paths[kind] = target
        except Exception as error:
            target.unlink(missing_ok=True)
            result.warnings.append(f'A 版可編輯圖稿未產生（{kind}，{type(error).__name__}: {error}）')

    power_a = getattr(result, 'power_tree_a', None) or {}
    if power_a.get('edges') is not None:
        def power():
            from autocheck.report.power_tree_diagram import layout
            from autocheck.report.power_tree_drawio import render_drawio
            boxes = layout(power_a, None, result.a.label, result.a.label, show_internal=False)
            boxes_all = layout(power_a, None, result.a.label, result.a.label, show_internal=True)
            return render_drawio(boxes, f'Power tree 方框圖 {result.a.label}　產生 {result.generated}',
                                 a_label=result.a.label, b_label=result.a.label, diff=False, boxes_all=boxes_all)

        write_a('power_drawio_a', f'power_tree_{result.a.label}_A.drawio', power)
    i2c_a = getattr(result, 'i2c_map_a', None) or {}
    if i2c_a.get('buses'):
        def i2c_diagram():
            from autocheck.report.i2c_block_drawio import render_drawio
            return render_drawio(i2c_a, f'I2C map 陽春版 {result.a.label}　產生 {result.generated}')
        write_a('i2c_drawio_a', f'i2c_block_{result.a.label}_A.drawio', i2c_diagram)


def _safe_workbook(workbook):
    from openpyxl.cell.cell import Cell
    from openpyxl.formula import Tokenizer
    from openpyxl.utils import get_column_letter
    from autocheck.report.xlsx_tracker import BASE_URL_CELL, COL
    # Column letters follow the workbook layout the engine writes (2026-10-27 inserted 父項 after 狀態), never a
    # hard-coded letter, so an engine update cannot silently turn the PDF links into plain text.
    dest_col, rule_col = get_column_letter(COL['PDF 目的地']), get_column_letter(COL['規則'])
    base_ref = BASE_URL_CELL.replace('$', '')
    base = workbook['版本資訊'][base_ref]
    if not isinstance(base.value, str) or not re.fullmatch(r'[^/\\:\x00-\x1f]{1,240}\.pdf', base.value):
        base.value = ''
    base.data_type = 's'
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if not isinstance(cell, Cell):
                    continue
                cell.hyperlink = None
                if cell.data_type != 'f':
                    continue
                trusted = False
                if sheet.title == '統計' and 2 <= cell.column <= 6:
                    try:
                        tokens = Tokenizer(cell.value).items
                        trusted = bool(tokens) and all(
                            (t.type == 'FUNC' and (t.subtype == 'CLOSE' or t.value in ('COUNTIF(', 'COUNTIFS(', 'COUNTBLANK(', 'COUNTA(', 'SUM(')))
                            or (t.type == 'OPERAND' and (t.subtype in ('TEXT', 'NUMBER') or (t.subtype == 'RANGE' and re.fullmatch(r'(?:追蹤表!)?\$?[A-Z]{1,2}\$?[1-9]\d*(?::\$?[A-Z]{1,2}\$?[1-9]\d*)?', t.value))))
                            or (t.type == 'SEP' and t.subtype == 'ARG') or (t.type == 'OP_IN' and t.value == '+')
                            for t in tokens)
                    except Exception:
                        pass
                elif sheet.title == '規則說明' and cell.column == 6:
                    trusted = bool(re.fullmatch(r'=COUNTIF\(追蹤表!\$' + rule_col + r'\$2:\$' + rule_col + r'\$[1-9]\d*,A[1-9]\d*\)', cell.value))
                elif sheet.title == '追蹤表' and cell.column == COL['PDF']:
                    r = cell.row
                    trusted = cell.value == (f'=IF(OR(版本資訊!{BASE_URL_CELL}="",{dest_col}{r}=""),"",'
                                             f'HYPERLINK(版本資訊!{BASE_URL_CELL}&"#nameddest="&{dest_col}{r},"PDF"))')
                elif sheet.title.startswith('Power tree 圖 ('):
                    cell_ref = re.escape(BASE_URL_CELL)
                    trusted = bool(re.fullmatch(r'=IF\(版本資訊!' + cell_ref + r'="","((?:[^"]|"")*)",HYPERLINK\(版本資訊!' + cell_ref + r'&"#page=[1-9]\d*","\1"\)\)', cell.value))
                if not trusted:
                    cell.data_type = 's'


def _run(payload, root, progress):
    a_label, b_label = _label(payload.get('aLabel'), 'A'), _label(payload.get('bLabel'), 'B')
    a, b = _classified(payload['a'], a_label), _classified(payload['b'], b_label)
    options = payload.get('options') or {}
    custom = _yaml_doc(_option_text(options['config'])) if options.get('config') is not None else {}
    library = _library({'yaml': _option_text(options['library'])} if options.get('library') is not None else {}, root)
    if library['errors']:
        raise ValueError('使用者位址庫：' + '；'.join(library['errors']))
    custom.setdefault('i2c', {})
    if not isinstance(custom['i2c'], dict):
        raise ValueError('i2c 設定必須是欄位對照')
    custom['i2c'].update(library=str(i2c.LIB_DEFAULT), user_library=str(root / 'i2c_devices_user.yaml'),
                         datasheets_dir=str(root / 'datasheets'))
    custom.setdefault('power_tree', {})
    if not isinstance(custom['power_tree'], dict):
        raise ValueError('power_tree 設定必須是欄位對照')
    custom['power_tree']['library'] = str(DEFAULT.parent / 'power_parts.yaml')
    custom['publish'] = {'base_url': '', 'static_dir': 'static'}
    config = root / 'effective.yaml'
    config.write_text(yaml.safe_dump(custom, allow_unicode=True), encoding='utf-8')
    option = lambda key: options.get(key, {}).get('file') if isinstance(options.get(key), dict) else None
    os.environ['AUTOCHECK_REVIEW_DB'] = str(root / 'review.db')
    progress('解析 A/B netlist、BOM、DSN 並執行原始檢查')
    result = run(a, b, a_label, b_label, config, option('railMap'), option('railDomain'),
                 option('i2cDeclared'), root / 'datasheets')
    result.warnings.extend(library['warnings'])
    rows = expand_rows(result.findings, result.cap_rows, result.b)
    # 2026-10-27 B-full: consequence rows of a NET_SHORTED / CAP_ONLY_NET root cause point at their parent, exactly as
    # cli.write_outputs links the rows it writes to Excel, so the website rows and the workbook agree.
    link_root_causes(rows, result.cfg)
    out = root / 'out'
    progress('產生原始 PDF、Excel、JSON、CSV、圖表與可編輯 draw.io')
    paths = {}
    failed_schematic = None
    try:
        paths = write_outputs(result, out, schematic_pdf=None, base_url='/__autocheck_local__', publish=False)
    except Exception as error:
        result.warnings.append(f'報表輸出未完整完成（{type(error).__name__}: {error}），僅提供已成功產生的檔案')
        trace = error.__traceback__
        annotation_failed = False
        while trace:
            annotation_failed |= trace.tb_frame.f_code.co_name == 'annotate_pdf'
            trace = trace.tb_next
        if annotation_failed:
            failed_schematic = out / f'AutoCheck_{a_label}_vs_{b_label}_schematic_{b_label}.pdf'
            pdf = result.files['b'].pop('pdf', None)
            try:
                paths = write_outputs(result, out, schematic_pdf=None, base_url=None, publish=False)
            except Exception as fallback_error:
                result.warnings.append(f'其餘報表輸出失敗（{type(fallback_error).__name__}: {fallback_error}）')
            finally:
                result.files['b']['pdf'] = pdf
    _drawio_outputs(result, out, paths)
    schematic = next((p for p in out.glob('*.pdf') if '_schematic_' in p.name and p != failed_schematic and p.stat().st_size), None)
    if schematic:
        prefix = f'/__autocheck_local__/{a_label}_vs_{b_label}/{schematic.name}'
        for path in out.glob('*.html'):
            path.write_text(path.read_text(encoding='utf-8').replace(prefix, schematic.name), encoding='utf-8')
        for path in out.glob('*.drawio'):
            path.write_text(path.read_text(encoding='utf-8').replace(escape(prefix, quote=True),
                                                                   escape(schematic.name, quote=True)), encoding='utf-8')
    for path in out.glob('*.xlsx'):
        from openpyxl import load_workbook
        workbook = load_workbook(path, keep_links=False)
        if schematic:
            workbook['版本資訊']['B5'] = schematic.name
        _safe_workbook(workbook)
        workbook.save(path)
        review.store_xlsx(root / 'review.db', f'{a_label}_vs_{b_label}', path)
    expected = {'pdf': f'AutoCheck_{a_label}_vs_{b_label}.pdf', 'xlsx': f'AutoCheck_tracker_{a_label}_vs_{b_label}.xlsx',
                'json': f'findings_{a_label}_vs_{b_label}.json', 'caps': f'cap_table_{b_label}.csv',
                'todo': 'rail_map_todo.csv', 'domains': f'rail_domain_{b_label}.csv',
                'circuit_b': f'circuit_{b_label}.txt', 'circuit_a': f'circuit_{a_label}.txt',
                'ptree_html': f'power_tree_{b_label}.html', 'i2c_csv': f'i2c_map_{b_label}.csv',
                'i2c_html': f'i2c_map_{b_label}.html', 'i2c_block_html': f'i2c_block_{b_label}.html',
                'power_drawio_b': f'power_tree_{b_label}.drawio', 'i2c_drawio_b': f'i2c_block_{b_label}.drawio'}
    known = {str(out / name): key for key, name in expected.items()}
    known.update({str(path): key for key, path in paths.items()})
    # The 2026-10-20+ engine writes a sidecar per annotated schematic (B and A); keys must stay unique because the
    # page uses them as React keys and looks artifacts up by key.
    known[str(out / f'AutoCheck_{a_label}_vs_{b_label}_schematic_{b_label}.annotations.json')] = 'annotations'
    known[str(out / f'AutoCheck_{a_label}_vs_{b_label}_schematic_{a_label}.annotations.json')] = 'annotations_a'
    artifacts = []
    used = set()
    for path in sorted(out.iterdir()):
        if not path.is_file() or not path.stat().st_size or path == failed_schematic:
            continue
        if failed_schematic and path == failed_schematic.with_suffix('.annotations.json'):
            continue
        key = known.get(str(path)) or ('annotations' if path.name.endswith('.annotations.json') else path.stem)
        unique, n = key, 2
        while unique in used:
            unique, n = f'{key}_{n}', n + 1
        used.add(unique)
        artifacts.append(_artifact(path, unique))
    return {'meta': meta(result), 'findings': [f.to_dict() for f in result.findings],
            'rows': [dict(dataclasses.asdict(row), dest=row.dest) for row in rows], 'overview': result.overview,
            'capRows': result.cap_rows, 'railTodo': result.rail_todo, 'railDomain': result.rail_domain,
            'powerTree': result.power_tree, 'powerTreeA': result.power_tree_a,
            'i2cMap': result.i2c_map, 'i2cMapA': result.i2c_map_a,
            'inputIssues': result.input_issues, 'warnings': result.warnings,
            'library': library['entries'], 'artifacts': artifacts,
            'rootcause': getattr(result, 'rootcause', None) or {}}


def _tracker(payload, root):
    from openpyxl import load_workbook
    a_label, b_label = _label(payload.get('aLabel'), 'A'), _label(payload.get('bLabel'), 'B')
    rows = [TrackRow(**{key: value for key, value in row.items() if key in TrackRow.__dataclass_fields__}) for row in payload['rows']]
    replies = payload.get('replies') or {}
    if not isinstance(replies, dict):
        raise ValueError('追蹤回覆必須是欄位對照')
    fields = {'disp': '處置', 'owner': '負責人', 'target_rev': '預計版次', 'note': '備註', 'filled': '填寫日期'}
    converted = {}
    for key, values in replies.items():
        if not isinstance(values, dict):
            raise ValueError('追蹤回覆格式無效')
        converted[key] = {}
        for source, target in fields.items():
            value = values.get(source) or ''
            if not isinstance(value, str) or len(value) > 10000:
                raise ValueError('回覆欄位必須是文字且不得超過 10000 字')
            converted[key][target] = value
    workbook = load_workbook(payload['template'], keep_links=False)
    review.fill_replies(workbook['追蹤表'], converted)
    from autocheck.report.xlsx_tracker import COL
    for row in workbook['追蹤表'].iter_rows(min_row=2):
        for field in fields.values():
            cell = row[COL[field] - 1]
            if cell.value is not None:
                cell.data_type = 's'
    _safe_workbook(workbook)
    out = root / f'AutoCheck_tracker_{a_label}_vs_{b_label}.xlsx'
    workbook.save(out)
    return _artifact(out, 'xlsx')


def autocheck_dispatch(request_json, progress):
    request = json.loads(request_json)
    kind, payload = request['type'], request['payload']
    root = Path(payload['root'])
    if kind == 'run':
        result = _run(payload, root, progress)
    elif kind == 'tracker':
        result = _tracker(payload, root)
    elif kind == 'library':
        result = _library(payload, root)
    elif kind == 'datasheet':
        progress('從原始 PDF 文字／表格擷取候選位址規則')
        result = draft_from_pdf(payload['file'], str(payload.get('name') or ''), str(payload.get('mpn') or ''))
        result['entry_yaml'] = entry_yaml(result['entry'])
        result['errors'], result['warnings'] = validate_entry(result['entry'])
        result['pageTexts'] = [{'page': page, 'text': page_text(payload['file'], page)} for page in result['pages'][:4]]
    else:
        raise ValueError('不支援的 AutoCheck 工作')
    return json.dumps(_plain(result), ensure_ascii=False, allow_nan=False)
